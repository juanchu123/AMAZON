"""
prediccion.py — lo que el agente sabe de cada producto: ticket medio, conversión, el modelo
de keyword_ml.py y las keywords candidatas para rellenar huecos (§2.2, §2.3).

- P(compra|clic) de una keyword = modelo de keyword_ml.py (si el producto tiene perfil y ≥ 10
  compras en el histórico) o, si no, la conversión media del producto. Después se mezcla con
  los clics MADUROS propios de la keyword (cuantos más tenga, más pesan sus datos).
- "ACOS predicho" de una candidata = CPC para competir / (P(compra|clic) × ticket).
  CPC para competir = la "puja recomendada baja" que da Amazon para esa frase (lo mínimo para
  salir con cierta regularidad); si la frase es nueva, la mediana de las frases parecidas del
  producto (las específicas por un lado, las genéricas —mucho más caras— por otro); si no hay
  ninguna, su CPC medio histórico.
- Escalonado por coincidencia (Amazon Ads Academy; histórico propio: la amplia convierte la mitad):
  P(compra|clic) de una frase nunca es mayor en Amplia que en Frase, ni en Frase que en Exacta
  (la amplia caza lo mismo que la frase y más cosas). Se fuerza con una regresión isotónica
  ponderada por los clics de cada coincidencia; así las pujas quedan amplia ≤ frase ≤ exacta.
  Sin modelo fiable, la conversión media del producto se reparte por coincidencia con lo que
  convierte cada una en el histórico de toda la cuenta.
"""

import statistics
from collections import Counter
from dataclasses import dataclass, field

import config
import keyword_ml as kml
import validacion
from documento import num

PERFIL_DE_ASIN = {p["asin"]: n for n, p in kml.PERFILES.items()}


@dataclass
class Producto:
    asin: str
    nombre: str = ""
    sku: str = ""
    perfil: str | None = None
    clics: float = 0.0
    coste: float = 0.0
    compras: float = 0.0
    ventas: float = 0.0
    cpc_competir: float | None = None               # mediana de la "puja rec. baja" de todas sus keywords
    cpc_especificas: float | None = None            # … de las que llevan su palabra distintiva
    cpc_genericas: float | None = None              # … de las genéricas (suelen ser mucho más caras)
    factor_cpc: float = 1.0                         # CPC pagado / puja rec. baja (≥ 1), config.FACTOR_CPC_*
    filas: list = field(default_factory=list)       # keywords del histórico (formato keyword_ml)
    clics_coincidencia: dict = field(default_factory=dict)   # clics del histórico por coincidencia
    modelo: tuple | None = None                     # (vec, modelo) de keyword_ml

    @property
    def ticket(self):
        return self.ventas / self.compras if self.compras else None

    @property
    def conversion(self):
        return self.compras / self.clics if self.clics else None

    @property
    def modelo_fiable(self):
        return self.modelo is not None and self.compras >= config.MIN_COMPRAS_MODELO

    @property
    def corto(self):
        return config.NOMBRES_CORTOS.get(self.asin) or (self.perfil or self.nombre.split(" ")[0] or self.asin).capitalize()


def isotonica(valores, pesos):
    """Regresión isotónica creciente (pool adjacent violators) ponderada."""
    bloques = [[v, w, 1] for v, w in zip(valores, pesos)]
    i = 0
    while i < len(bloques) - 1:
        if bloques[i][0] > bloques[i + 1][0] + 1e-12:
            (v1, w1, n1), (v2, w2, n2) = bloques[i], bloques[i + 1]
            bloques[i] = [(v1 * w1 + v2 * w2) / (w1 + w2), w1 + w2, n1 + n2]
            del bloques[i + 1]
            i = max(i - 1, 0)
        else:
            i += 1
    return [v for v, _, n in bloques for _ in range(n)]


class Catalogo:
    def __init__(self, ruta_historico=None):
        self.productos = {}
        self._cache_p = {}
        self.conv_coincidencia = {}     # conversión de cada coincidencia en todo el histórico
        self.clics_coincidencia = {}
        ruta = ruta_historico or config.HISTORICO_XLSX
        if ruta and ruta.exists():
            self._cargar(ruta)
        clics, compras = Counter(), Counter()
        for p in self.productos.values():
            for f in p.filas:
                clics[f["coincidencia"]] += f["clics"]
                compras[f["coincidencia"]] += f["compras"]
        self.clics_coincidencia = dict(clics)
        self.conv_coincidencia = {m: compras[m] / clics[m] for m in kml.COINCIDENCIAS if clics[m]}

    def _cargar(self, ruta):
        from openpyxl import load_workbook
        wb = load_workbook(ruta, read_only=True, data_only=True)
        if "Resumen por producto" in wb.sheetnames:
            it = wb["Resumen por producto"].iter_rows(values_only=True)
            cab = [str(c or "") for c in next(it)]
            for r in it:
                d = dict(zip(cab, r))
                asin = str(d.get("ASIN") or "")
                if not asin.startswith("B0"):
                    continue
                self.productos[asin] = Producto(
                    asin=asin, nombre=str(d.get("Producto") or ""), sku=str(d.get("SKU") or ""),
                    perfil=PERFIL_DE_ASIN.get(asin), clics=num(d.get("Clics")), coste=num(d.get("Coste (€)")),
                    compras=num(d.get("Compras")), ventas=num(d.get("Ventas (€)")),
                    cpc_competir=num(d.get("CPC (€)"), None))
        for asin, p in self.productos.items():
            try:
                if p.perfil:
                    kml.usar_perfil(p.perfil)
                filas, _, _, _, _ = kml.cargar_excel(str(ruta), asin=asin)
            except SystemExit:
                continue
            p.filas = filas
            p.clics_coincidencia = dict(sum((Counter({f["coincidencia"]: f["clics"]}) for f in filas), Counter()))
            rec = [f for f in filas if f.get("puja_rec_baja", 0) > 0]
            if rec:
                p.cpc_competir = statistics.median(f["puja_rec_baja"] for f in rec)
                esp = [f["puja_rec_baja"] for f in rec if self._especifica(p, f["keyword"])]
                gen = [f["puja_rec_baja"] for f in rec if not self._especifica(p, f["keyword"])]
                p.cpc_especificas = statistics.median(esp) if esp else None
                p.cpc_genericas = statistics.median(gen) if gen else None
            ratios = [f["coste"] / f["clics"] / f["puja_rec_baja"] for f in filas
                      if f["clics"] >= config.FACTOR_CPC_MIN_CLICS and f.get("puja_rec_baja", 0) > 0]
            if len(ratios) >= config.FACTOR_CPC_MIN_FILAS:
                p.factor_cpc = max(1.0, statistics.median(ratios))
            if p.perfil and sum(f["compras"] for f in filas) > 0:
                p.modelo = kml.entrenar(filas, "compras", "clics")

    @staticmethod
    def _especifica(p, texto):
        return bool(p.perfil) and any(w in kml.PERFILES[p.perfil]["especificas"] for w in kml.tokens_contenido(texto))

    def cpc_para(self, asin, texto, rec_propia=None):
        """CPC que se pagará por una frase: la puja sugerida por Amazon para ELLA si existe (si no, la
        mediana de las frases parecidas del producto, específicas o genéricas), por el factor de
        calibración del producto (lo que de verdad se ha pagado frente a lo que sugería Amazon)."""
        p = self.producto(asin)
        if rec_propia and rec_propia > 0:
            return rec_propia * p.factor_cpc
        propia = p.cpc_especificas if self._especifica(p, texto) else p.cpc_genericas
        base = propia or p.cpc_competir
        return base * p.factor_cpc if base else base

    def producto(self, asin):
        if asin not in self.productos:
            self.productos[asin] = Producto(asin=asin, perfil=PERFIL_DE_ASIN.get(asin))
        return self.productos[asin]

    # ------------------------------------------------------------ predicciones
    def _por_coincidencia_sin_modelo(self, p):
        """Conversión media del producto repartida por coincidencia según el histórico de la cuenta,
        de forma que la media ponderada por los clics del producto siga siendo su conversión."""
        base = p.conversion or 0.0
        conv = self.conv_coincidencia
        if not base or len(conv) < len(kml.COINCIDENCIAS):
            return [base] * len(kml.COINCIDENCIAS)
        total = sum(self.clics_coincidencia.values()) or 1.0
        media = sum(self.clics_coincidencia[m] * conv[m] for m in kml.COINCIDENCIAS) / total
        mezcla = p.clics_coincidencia if sum(p.clics_coincidencia.values()) else self.clics_coincidencia
        n = sum(mezcla.get(m, 0) for m in kml.COINCIDENCIAS) or 1.0
        media_mezcla = sum(mezcla.get(m, 0) * conv[m] / media for m in kml.COINCIDENCIAS) / n
        return [base * conv[m] / media / (media_mezcla or 1.0) for m in kml.COINCIDENCIAS]

    def p_modelo(self, asin, texto, coincidencia):
        """P(compra|clic) a priori (sin los datos propios del elemento), escalonada por coincidencia."""
        p = self.producto(asin)
        if coincidencia not in kml.COINCIDENCIAS:
            return p.conversion or 0.0
        clave = (asin, str(texto).strip().lower())
        if clave not in self._cache_p:
            if p.modelo_fiable:
                kml.usar_perfil(p.perfil)
                crudos = [kml.predecir(*p.modelo, texto, m) for m in kml.COINCIDENCIAS]
                pesos = [max(1.0, p.clics_coincidencia.get(m, 0)) for m in kml.COINCIDENCIAS]
            else:
                crudos = self._por_coincidencia_sin_modelo(p)
                pesos = [max(1.0, self.clics_coincidencia.get(m, 0)) for m in kml.COINCIDENCIAS]
            self._cache_p[clave] = isotonica(crudos, pesos)
        return self._cache_p[clave][kml.COINCIDENCIAS.index(coincidencia)]

    def p_compra(self, asin, texto, coincidencia, maduro):
        """Mezcla el modelo con los clics maduros propios: (compras + K·p) / (clics + K)."""
        p0 = self.p_modelo(asin, texto, coincidencia)
        k = config.PRIOR_CLICS
        return (maduro.compras + k * p0) / (maduro.clics + k)

    # ------------------------------------------------------------ candidatas (§2.3)
    def _ajenas(self, p):
        """Palabras que identifican a OTROS productos: nunca en las keywords de este."""
        ajenas = set().union(*(q["especificas"] for n, q in kml.PERFILES.items() if n != p.perfil))
        return ajenas - (kml.PERFILES[p.perfil]["especificas"] if p.perfil else set())

    def evaluar(self, asin, texto, coincidencia=None, pc=None, rec=None):
        """Puntúa una frase para un producto con el mismo filtro que usa el agente para decidir.
        Devuelve {p, cpc, acos_pred, puja, pasa, motivo}. pc: P(compra|clic) si ya se conoce
        (si no, la del modelo); rec: puja sugerida por Amazon para esa frase, si se sabe."""
        coincidencia = coincidencia or config.COINCIDENCIA_NUEVAS
        p = self.producto(asin)
        if p.perfil:
            kml.usar_perfil(p.perfil)
        pc = self.p_modelo(asin, texto, coincidencia) if pc is None else pc
        cpc = self.cpc_para(asin, texto, rec)
        res = {"p": pc, "cpc": cpc, "acos_pred": None, "puja": None, "pasa": False}
        valida, motivo = validacion.keyword(texto)
        if not valida:
            return res | {"motivo": f"Amazon no la aceptaría: {motivo}"}
        otras = sorted(kml.firma(texto) & self._ajenas(p))
        if otras:
            return res | {"motivo": "lleva palabras de otro producto: " + ", ".join(otras)}
        if not p.ticket or not cpc or pc <= 0:
            return res | {"motivo": "sin datos del producto para predecir"}
        acos_pred = cpc / (pc * p.ticket)
        puja = pc * p.ticket * config.ACOS_MAX_KEYWORD_NUEVA
        if not p.modelo_fiable:
            puja = min(puja, config.PUJA_MAX_SIN_MODELO)
        pasa = acos_pred <= config.ACOS_MAX_KEYWORD_NUEVA
        return res | {"acos_pred": acos_pred, "puja": max(config.PUJA_MINIMA_AMAZON, round(puja, 2)), "pasa": pasa,
                      "motivo": "entra si hay hueco en un grupo del producto" if pasa else
                                f"ACOS predicho {acos_pred:.0%} > {config.ACOS_MAX_KEYWORD_NUEVA:.0%}"}

    def candidatas(self, asin, ya_usadas, investigacion=()):
        """Keywords candidatas para el producto que pasan el filtro de ACOS predicho ≤ 30 %,
        ordenadas de mejor a peor. ya_usadas: firmas (keyword_ml.firma) que el producto ya tiene
        en cualquier campaña o estado (nunca se duplica ni se reactiva una pausada).
        Devuelve [{texto, coincidencia, p, acos_pred, puja, fuente, motivo}]."""
        p = self.producto(asin)
        ticket = p.ticket
        if not ticket:
            return []
        if p.perfil:
            kml.usar_perfil(p.perfil)
        vistas, brutas = set(ya_usadas), []

        # 1) keywords del histórico del producto que ya vendieron (nunca Amplia genérica)
        for f in sorted(p.filas, key=lambda f: -f["compras"]):
            if f["compras"] < 1:
                continue
            toks = kml.tokens_contenido(f["keyword"])
            especifica = p.perfil and any(w in kml.PERFILES[p.perfil]["especificas"] for w in toks)
            if f["coincidencia"] == "Amplia" and not especifica:
                continue
            k = config.PRIOR_CLICS
            pc = (f["compras"] + k * self.p_modelo(asin, f["keyword"], f["coincidencia"])) / (f["clics"] + k)
            brutas.append((f["keyword"], f["coincidencia"], pc, "histórico",
                           f"Histórico: {int(f['clics'])} clics, {int(f['compras'])} compras", f.get("puja_rec_baja")))
        # 2) investigación de mercado (LLM) — solo propone, el filtro decide
        for fila in investigacion:
            texto = str(fila.get("Palabra clave") or "").strip().lower()
            if texto:
                rec = num(fila.get("Puja sugerida (€)"), None)
                brutas.append((texto, config.COINCIDENCIA_NUEVAS, self.p_modelo(asin, texto, config.COINCIDENCIA_NUEVAS),
                               "investigación", str(fila.get("Motivo") or "")[:200], rec if rec and rec > 0 else None))
        # 3) frases nuevas que genera el modelo de keyword_ml
        if p.modelo_fiable:
            try:
                recs, _ = kml.recomendar(p.filas, p.nombre, set(), top=40, acos_objetivo=100 * config.ACOS_MAX_KEYWORD_NUEVA,
                                         coincidencia=config.COINCIDENCIA_NUEVAS)
                for r in recs:
                    brutas.append((r["palabra_clave"], config.COINCIDENCIA_NUEVAS, r["prob_compra_por_clic"], "modelo",
                                   r["motivo"] or "Frase nueva del modelo", None))
            except Exception:
                pass

        out = []
        for texto, coinc, pc, fuente, motivo, rec in brutas:
            fi = kml.firma(texto)
            if not fi or fi in vistas:
                continue
            ev = self.evaluar(asin, texto, coinc, pc, rec)
            if ev["acos_pred"] is not None:
                vistas.add(fi)
            if ev["pasa"]:
                out.append({"texto": texto, "coincidencia": coinc, "p": pc, "acos_pred": ev["acos_pred"],
                            "puja": ev["puja"], "fuente": fuente, "motivo": motivo})
        out.sort(key=lambda c: c["acos_pred"])
        return out

    def candidatas_asin(self, asin, competidores, ya_usados):
        """ASIN de competencia (hoja Competencia, confirmados por Juan) para grupos de pruebas.
        Conversión = la media del producto (aún no hay datos por ASIN competidor)."""
        p = self.producto(asin)
        if not p.ticket or not p.cpc_competir or not p.conversion:
            return []
        acos_pred = p.cpc_competir / (p.conversion * p.ticket)
        if acos_pred > config.ACOS_MAX_KEYWORD_NUEVA:
            return []
        puja = round(max(config.PUJA_MINIMA_AMAZON, p.conversion * p.ticket * config.ACOS_MAX_KEYWORD_NUEVA), 2)
        return [{"texto": str(c["ASIN competidor"]).strip().upper(), "coincidencia": "ASIN", "p": p.conversion,
                 "acos_pred": acos_pred, "puja": puja, "fuente": "competencia", "motivo": str(c.get("Motivo") or "")}
                for c in competidores if str(c["ASIN competidor"]).strip().upper() not in ya_usados]
