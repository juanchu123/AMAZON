"""
crear_memoria.py — genera los dos archivos para lanzar y seguir las campañas:

  1. resultados/memoria.xlsx          -> la MEMORIA: configuración, segmentación (como Amazon),
                                          seguimiento cada 3 días, un ticket por cambio y el histórico.
  2. resultados/bulk_AAAA-MM-DD.xlsx  -> hoja masiva para subir en Amazon Ads > Operaciones en bloque
                                          (rellena la plantilla oficial de Juan). Crea todo EN PAUSA.

Los dos salen de los mismos datos, así que siempre coinciden. Campañas:
  - Pinza - Principal V2  : 5 keywords históricas + 5 especiales del modelo (keyword_ml.py)
  - Pinza - Pruebas       : ASIN de competencia (+ categoría, que se añade a mano)
  - Rejilla - Principal V2: keywords históricas + especiales del modelo, perfil 'rejilla'

Uso:  python crear_memoria.py [--fecha 2026-09-26] [--plantilla ruta/AdvertisingBulksheetTemplate.xlsx]
Después hay que recalcular fórmulas (abrir en Excel, o scripts/recalc.py de LibreOffice).
"""

import argparse
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import keyword_ml as k

EXCEL_HIST = "FreshFinder_Amazon_Ads_historico.xlsx"
ACOS_OBJ = 0.35
PRESUPUESTO_DIARIO = 7.0          # pedido por Juan (26/09/2026): 7 €/día en cada campaña
CARTERA_MENSUAL = 100.0           # tope duro de CLAUDE.md; Juan decide si lo sube
PUJA_PRUEBA = 0.30

NEG_COMUNES = [
    ("ventosa", "Otro tipo de montaje (tienes otro producto con ventosa)."),
    ("magnetico", "Soportes magnéticos: otro producto."),
    ("iman", "Soportes de imán: otro producto."),
    ("parabrisas", "No va al parabrisas."),
    ("moto", "Soportes de moto: otro producto."),
    ("bicicleta", "Soportes de bici: otro producto."),
    ("cargador", "Soportes con carga inalámbrica: otro producto y otro precio."),
]

CAMPANAS = [
    {
        "nombre": "Pinza - Principal V2", "grupo": "Pinza - Principal V2", "perfil": "pinza",
        "sku": "O8-5W7J-DSK1", "tipo": "Palabras clave", "puja_grupo": 0.40,
        "n_hist": 5,
        "especiales": [  # elegidas en conversaciones anteriores (top del modelo + criterio)
            ("soporte móvil coche pinza 360", "Top del modelo; todas sus palabras tienen datos."),
            ("soporte móvil coche pinza para espejo retrovisor", "Retrovisor: sitio de montaje del título sin cubrir."),
            ("soporte móvil coche pinza para parasol", "Parasol: sitio de montaje del título sin cubrir."),
            ("soporte móvil coche pinza ajustable", "Siguiente del ranking; 'ajustable' es la primera característica del título."),
            ("soporte móvil coche pinza para iphone", "Búsquedas por modelo de móvil; el título dice 'compatible con iPhone'."),
        ],
        "negativas": NEG_COMUNES + [("rejilla", "Es el soporte de rejilla (B0DHYBY6MS), tiene su propia campaña.")],
    },
    {
        "nombre": "Pinza - Pruebas", "grupo": "Pruebas - Competencia", "perfil": "pinza",
        "sku": "O8-5W7J-DSK1", "tipo": "Productos", "puja_grupo": PUJA_PRUEBA,
        "asins": [  # sacados de búsqueda web: verificar en amazon.es antes de activar
            ("B07H996QMD", "OcioDual — pinza para salpicadero. Competidor directo."),
            ("B0BS1TQJH2", "Soporte con pinza para salpicadero (gravedad). Competidor directo."),
            ("B08N4KR6PK", "BEENLE — salpicadero + espejo retrovisor."),
            ("B09P3SZB6P", "AUTOZOCO — pinza para salpicadero con resorte."),
        ],
        "categoria": ('categoría="Soportes montados en salpicadero para automóviles" precio>12€ estrellas≤4',
                      "Categoría de más vendidos donde compite la pinza, filtrada a fichas más caras o peor valoradas. SE AÑADE A MANO (la hoja masiva pide el código numérico de categoría)."),
        "negativos_producto": [("B0DCZS1NR6", "Tu propia ficha: no pagar por salir en ella.")],
    },
    {
        "nombre": "Rejilla - Principal V2", "grupo": "Rejilla - Principal V2", "perfil": "rejilla",
        "sku": "5E-I8NY-S191", "tipo": "Palabras clave", "puja_grupo": 0.40,
        "n_hist": 5, "n_total": 10,       # si no hay 5 históricas válidas, se completa con especiales
        "especiales": None,               # None = las mejores del modelo (keyword_ml.recomendar)
        "negativas": NEG_COMUNES + [
            ("pinza", "Es el soporte de pinza, tiene su propia campaña (la rejilla vendió con 'pinza' pero a ACOS 85%)."),
            ("salpicadero", "La rejilla no va al salpicadero."),
        ],
    },
]


# ================================================================== datos
def ajuste_historico(a, ticket, conv_media, min_clics=10):
    """Puja nueva de una keyword histórica: tabla de la sección 5.1 de marketingV2.md,
    sin pasar de su puja máxima rentable (su conversión real × su ticket × 35%).
    Con menos de min_clics su conversión es ruido: se usa la media del producto."""
    acos = a["coste"] / a["ventas"]
    pocos = a["clics"] < min_clics
    conv = conv_media if pocos else a["compras"] / a["clics"]
    pmax = conv * (ticket if pocos else a["ventas"] / a["compras"]) * ACOS_OBJ
    if acos < 0.30:
        nueva, regla = a["puja"] * 1.25, f"ACOS {acos:.1%} < 30% → subir 25%"
    elif acos <= 0.35:
        nueva, regla = a["puja"], f"ACOS {acos:.1%} en objetivo → se mantiene"
    elif acos <= 0.40:
        nueva, regla = a["puja"] * 0.90, f"ACOS {acos:.1%} algo alto → bajar 10%"
    else:
        nueva, regla = a["puja"] * 0.75, f"ACOS {acos:.1%} > 40% → bajar 25%"
    if nueva > pmax:
        nueva, regla = pmax, regla + f" (limitada a la puja máx. rentable {pmax:.2f} €" + \
            (f", con la conversión media del producto porque solo tiene {a['clics']:.0f} clics)" if pocos else ")")
    return round(nueva, 2), regla, pmax


def preparar():
    productos, lineas, historico = {}, [], []
    otras_especificas = lambda perfil: set().union(*(p["especificas"] for n, p in k.PERFILES.items() if n != perfil))
    for c in CAMPANAS:
        perfil = k.usar_perfil(c["perfil"])
        if c["perfil"] not in productos:
            filas, info, titulo, asin, existentes = k.cargar_excel(EXCEL_HIST, asin=perfil["asin"])
            vec, mod = k.entrenar(filas, "compras", "clics")
            ticket = sum(f["ventas"] for f in filas) / sum(f["compras"] for f in filas)
            conv = sum(f["compras"] for f in filas) / sum(f["clics"] for f in filas)
            productos[c["perfil"]] = dict(filas=filas, info=info, titulo=titulo, asin=asin, existentes=existentes,
                                          vec=vec, mod=mod, ticket=ticket, conv=conv, elegidas=set())
            for f in filas:
                historico.append((c["perfil"], asin, f))
        P = productos[c["perfil"]]
        c["asin"], c["titulo"], c["ticket"] = P["asin"], P["titulo"], P["ticket"]
        base = dict(campana=c["nombre"], grupo=c["grupo"], sku=c["sku"])

        if c["tipo"] == "Productos":
            for a, motivo in c["asins"]:
                lineas.append(dict(base, texto=f'asin="{a}"', match="Producto", origen="Prueba producto",
                                   puja=PUJA_PRUEBA, pmax=P["conv"] * P["ticket"] * ACOS_OBJ, motivo=motivo,
                                   ref="obj", antes=None))
            texto, motivo = c["categoria"]
            lineas.append(dict(base, texto=texto, match="Categoría", origen="Prueba categoría", puja=PUJA_PRUEBA,
                               pmax=P["conv"] * P["ticket"] * ACOS_OBJ, motivo=motivo, ref="obj", antes=None, manual=True))
            continue

        # históricas: se suma cada keyword + coincidencia de todos los grupos del producto
        agg = {}
        for f in P["filas"]:
            key = (k.normalizar(f["keyword"]), f["coincidencia"])
            a = agg.setdefault(key, dict(keyword=f["keyword"], coincidencia=f["coincidencia"], clics=0, coste=0,
                                         compras=0, ventas=0, puja=f["puja"], _cp=-1, grupos=[]))
            for x in ("clics", "coste", "compras", "ventas"):
                a[x] += f[x]
            a["grupos"].append(f["grupo"])
            if f["clics"] > a["_cp"]:
                a["puja"], a["_cp"] = f["puja"], f["clics"]
        otras = otras_especificas(c["perfil"])
        validas = []
        for a in sorted([a for a in agg.values() if a["compras"] > 0], key=lambda a: a["coste"] / a["ventas"]):
            toks = set(k.tokens_contenido(a["keyword"]))
            if toks & otras:
                continue  # palabra de OTRO producto (p. ej. "pinza" en la rejilla): va como negativa
            if a["coincidencia"] == "Amplia" and not toks & perfil["especificas"]:
                continue  # regla: nunca Amplia genérica (histórico: ACOS 122-177%)
            validas.append(a)
        for a in validas[:c["n_hist"]]:
            nueva, regla, pmax = ajuste_historico(a, P["ticket"], P["conv"])
            n_g = len(a["grupos"])
            motivo = (f"Histórico{' (suma de ' + str(n_g) + ' grupos)' if n_g > 1 else ''}: {a['clics']:.0f} clics, "
                      f"{a['compras']:.0f} compras, {a['coste']:.2f} € de gasto, {a['ventas']:.2f} € de ventas. {regla}.")
            lineas.append(dict(base, texto=a["keyword"], match=a["coincidencia"], origen="Histórica", puja=nueva,
                               pmax=pmax, motivo=motivo, ref=round(a["coste"] / a["ventas"], 4), antes=a["puja"]))
            P["elegidas"].add(k.firma(a["keyword"]))

        n_esp = c.get("n_total", 10) - len(validas[:c["n_hist"]])
        if c["especiales"] is not None:
            esp = c["especiales"][:n_esp]
        else:
            recs, _ = k.recomendar(P["filas"], P["titulo"], P["existentes"] | P["elegidas"], n_esp,
                                   acos_objetivo=ACOS_OBJ * 100, puja_minima=0.19)
            esp = [(r["palabra_clave"], "Del ranking del modelo. " + r["motivo"]) for r in recs]
        for kw, motivo in esp:
            p = k.predecir(P["vec"], P["mod"], kw, "Frase")
            pmax = p * P["ticket"] * ACOS_OBJ
            lineas.append(dict(base, texto=kw, match="Frase", origen="Especial IA", puja=round(pmax, 2), pmax=pmax,
                               motivo=f"{motivo} Modelo: {p:.1%} compra/clic → puja máx. rentable {pmax:.2f} € (ACOS 35%).",
                               ref="obj", antes=None))
    return productos, lineas, historico


# ================================================================== estilos
F = "Arial"
f_base, f_bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
f_head, f_title = Font(name=F, size=10, bold=True, color="FFFFFF"), Font(name=F, size=14, bold=True)
f_input, f_link = Font(name=F, size=10, color="0000FF"), Font(name=F, size=10, color="008000")
fill_head, fill_head2 = PatternFill("solid", fgColor="232F3E"), PatternFill("solid", fgColor="37475A")
fill_input = PatternFill("solid", fgColor="FFF2CC")
FILL_ORIGEN = {"Histórica": PatternFill("solid", fgColor="E2EFDA"), "Especial IA": PatternFill("solid", fgColor="DDEBF7"),
               "Prueba producto": PatternFill("solid", fgColor="FCE4D6"), "Prueba categoría": PatternFill("solid", fgColor="FCE4D6")}
red, green = PatternFill("solid", fgColor="F8CBAD"), PatternFill("solid", fgColor="C6EFCE")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap, center = Alignment(wrap_text=True, vertical="top"), Alignment(horizontal="center", vertical="center", wrap_text=True)
EUR, PCT, INT, DATE, X2 = '#,##0.00 "€";-#,##0.00 "€";-', '0.0%;-0.0%;-', '#,##0;-#,##0;-', "dd/mm/yyyy", "0.00;-0.00;-"


def header(ws, row, cols, fill=fill_head):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.alignment, cell.border = f_head, fill, center, border
    ws.row_dimensions[row].height = 42


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, r, c, v, fmt=None, font=f_base, fill=None, al=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.border = font, border
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if al:
        cell.alignment = al
    return cell


# ================================================================== memoria.xlsx
def crear_memoria(productos, lineas, historico, hoy, salida):
    wb = Workbook()
    # ---------------- Leyenda
    ws = wb.active
    ws.title = "Leyenda"
    widths(ws, [40, 16, 80])
    ws["A1"], ws["A1"].font = "Memoria de campañas FreshFinder (pinza + rejilla)", f_title
    ws["A2"], ws["A2"].font = f"Creada {hoy:%d/%m/%Y} con crear_memoria.py a partir de {EXCEL_HIST} y keyword_ml.py", f_base
    ws["A4"], ws["A4"].font = "PARÁMETROS", f_bold
    params = [  # fila 5..
        ("ACOS objetivo", ACOS_OBJ, PCT, "Confirmado por Juan."),
        ("Presupuesto diario por campaña", PRESUPUESTO_DIARIO, EUR, "Pedido por Juan (26/09/2026): 7 €/día en cada una de las 3 campañas."),
        ("Nº de campañas", len(CAMPANAS), "0", ""),
        ("Gasto máximo posible al mes", "=B6*B7*30", EUR, "21 €/día × 30 días. ⚠ Muy por encima del límite de 100 €/mes."),
        ("Tope mensual de la cartera", CARTERA_MENSUAL, EUR, "Límite duro de CLAUDE.md. Amazon para TODAS las campañas al llegar aquí. Súbelo solo si decides gastar más."),
        ("Días hasta agotar la cartera", "=IFERROR(B9/(B6*B7),0)", "0.0", "Con 21 €/día, la cartera de 100 € se agota en ~5 días."),
        ("Días entre revisiones de puja", 3, "0", "Pedido por Juan."),
        ("Rango de puja sobre la original", 0.50, PCT, "Pedido por Juan: ±50%."),
        ("Stop-loss: clics sin ventas", 15, "0", "Propuesto; pendiente de confirmar."),
        ("Clics mínimos para evaluar", 10, "0", "Con menos clics es ruido."),
        ("Días para evaluar un cambio", 7, "0", "Sponsored Products atribuye ventas hasta 7 días después del clic."),
    ]
    for i, (lab, val, fmt, nota) in enumerate(params, 5):
        put(ws, i, 1, lab)
        formula = isinstance(val, str) and val.startswith("=")
        put(ws, i, 2, val, fmt, f_base if formula else f_input, None if formula else fill_input)
        put(ws, i, 3, nota, al=wrap)
    ws["C8"].font = Font(name=F, size=10, bold=True, color="C00000")
    P_ACOS, P_RANGO, P_STOP, P_MIN, P_EVAL = "Leyenda!$B$5", "Leyenda!$B$12", "Leyenda!$B$13", "Leyenda!$B$14", "Leyenda!$B$15"
    r = 17
    ws.cell(row=r, column=1, value="CÓMO USAR LA MEMORIA").font = f_bold
    pasos = [
        f"1. Sube resultados/bulk_{hoy:%Y-%m-%d}.xlsx en Amazon Ads → Operaciones en bloque. Crea las 3 campañas EN PAUSA. Revisa el informe de resultados de Amazon.",
        "2. A mano: crea la cartera (T-001) y mete dentro las 3 campañas; añade la categoría de 'Pinza - Pruebas'. Activa las campañas cuando todo esté bien.",
        "3. En 'Tickets' pon la 'Fecha aplicado' de cada ticket y cambia su estado a 'Aplicado'.",
        "4. Cada 3 días: exporta Segmentación de cada campaña (rango DESDE LA CREACIÓN HASTA HOY) y copia una fila por keyword/ASIN en 'Seguimiento' (fecha, campaña, texto, coincidencia, puja, impresiones, clics, coste, compras, ventas).",
        "5. 'Segmentación' muestra el estado de cada keyword como en Amazon, con acción sugerida y stop-loss. Cada cambio que hagas = ticket nuevo (con la 'Base' copiada de Segmentación).",
        "6. Un ticket se evalúa a los 7 días y con ≥10 clics nuevos: 'Veredicto' dice Mejora/Empeora.",
        "7. Cada 2 semanas: descarga Términos de búsqueda (Amazon solo guarda 65 días) → negativas y keywords nuevas; y mira el informe de Ubicación.",
    ]
    for i, t in enumerate(pasos, r + 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font, c.alignment = f_base, wrap
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        ws.row_dimensions[i].height = 32
    r += len(pasos) + 2
    ws.cell(row=r, column=1, value="COLORES").font = f_bold
    for i, (fill, txt) in enumerate([(fill_input, "Celda para rellenar (texto azul)"), (FILL_ORIGEN["Histórica"], "Keyword histórica: ya vendió"),
                                     (FILL_ORIGEN["Especial IA"], "Keyword especial del modelo (mismas reglas que las históricas)"),
                                     (FILL_ORIGEN["Prueba producto"], "Campaña de pruebas: ASIN de competencia o categoría")], r + 1):
        put(ws, i, 1, "Ejemplo", fill=fill)
        ws.cell(row=i, column=2, value=txt).font = f_base
    r += 6
    ws.cell(row=r, column=1, value="AVISOS").font = f_bold
    for i, t in enumerate([
        "• Las campañas antiguas siguen en pausa: si se reactivan, sus keywords compiten con estas (tú contra ti).",
        "• Los anuncios solo salen con la OFERTA DESTACADA (Buy Box) y stock. 0 impresiones → mira eso antes de tocar pujas.",
        "• Los ASIN de competencia salen de una búsqueda web: comprueba en amazon.es que existen y son más caros o peor valorados que el tuyo.",
        "• Rejilla: el modelo casi no aprende de su histórico (mejora 0,4% sobre la media; casi todo eran búsquedas genéricas). Sus especiales van más a ciegas que las de la pinza.",
    ], r + 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font, c.alignment = f_base, wrap
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        ws.row_dimensions[i].height = 30

    # ---------------- Campañas
    wc = wb.create_sheet("Campañas")
    ccols = ["Campaña", "Grupo de anuncios", "Producto (ASIN)", "SKU", "Segmentación", "Presupuesto diario", "Estrategia de pujas",
             "Puja predet. grupo", "Ubicaciones", "Cartera", "Estado al crear", "Ticket medio (€)", "Nº segmentaciones", "Ticket"]
    widths(wc, [24, 24, 14, 15, 15, 12, 26, 11, 26, 26, 12, 12, 12, 9])
    wc["A1"], wc["A1"].font = "Configuración de campañas (tal cual se crean en Amazon Ads)", f_title
    header(wc, 3, ccols)
    for i, c in enumerate(CAMPANAS, 4):
        n = sum(1 for l in lineas if l["campana"] == c["nombre"])
        vals = [c["nombre"], c["grupo"], c["asin"], c["sku"], "Manual - " + c["tipo"], "=Leyenda!$B$6", "Pujas dinámicas: solo reducir",
                c["puja_grupo"], "0% en las 3 (revisar a las 2 semanas)", "FreshFinder - cartera", "Pausado", round(c["ticket"], 2), n, None]
        for j, v in enumerate(vals, 1):
            put(wc, i, j, v, EUR if j in (6, 8, 12) else None, f_link if j == 6 else f_base, al=wrap)
    # ---------------- Segmentación
    wsg = wb.create_sheet("Segmentación")
    scol = ["Campaña", "Estado", "Palabra clave / segmentación", "Tipo de coincidencia", "Puja (EUR)", "Impresiones", "Clics", "CTR",
            "Coste total (EUR)", "CPC (EUR)", "Compras", "Ventas (EUR)", "ACOS", "ROAS",
            "Origen", "Puja original", "Puja mín. -50%", "Puja máx. +50%", "Puja máx. rentable", "Ticket", "Última foto",
            "Stop-loss", "Acción sugerida", "Clave"]
    widths(wsg, [22, 10, 46, 12, 10, 11, 8, 8, 11, 9, 9, 11, 9, 8, 14, 10, 10, 10, 11, 8, 11, 24, 32, 0.1])
    wsg["A1"], wsg["A1"].font = "Segmentación — mismas columnas que el export de Amazon Ads. Se actualiza sola desde 'Seguimiento'.", f_bold
    header(wsg, 3, scol[:14])
    for i, c in enumerate(scol[14:], 15):
        cell = wsg.cell(row=3, column=i, value=c)
        cell.font, cell.fill, cell.alignment, cell.border = f_head, fill_head2, center, border
    S0 = 4
    SL = S0 + len(lineas) - 1
    for i, l in enumerate(lineas):
        r = S0 + i
        clave = f"X{r}"
        last = f"_xlfn.MAXIFS(Seguimiento!$A:$A,Seguimiento!$R:$R,{clave})"
        pull = lambda col: f"=IFERROR(SUMIFS(Seguimiento!${col}:${col},Seguimiento!$R:$R,{clave},Seguimiento!$A:$A,{last}),0)"
        vals = {
            1: l["campana"],
            2: f'=IFERROR(INDEX(Seguimiento!$E:$E,MATCH({clave}&"#"&{last},Seguimiento!$S:$S,0)),"Activado")',
            3: l["texto"], 4: l["match"],
            5: f"=IFERROR(IF({last}=0,P{r},SUMIFS(Seguimiento!$G:$G,Seguimiento!$R:$R,{clave},Seguimiento!$A:$A,{last})),P{r})",
            6: pull("H"), 7: pull("I"), 8: f"=IFERROR(G{r}/F{r},0)", 9: pull("J"), 10: f"=IFERROR(I{r}/G{r},0)",
            11: pull("K"), 12: pull("L"), 13: f"=IFERROR(I{r}/L{r},0)", 14: f"=IFERROR(L{r}/I{r},0)",
            15: l["origen"], 16: l["puja"], 17: f"=ROUND(P{r}*(1-{P_RANGO}),2)", 18: f"=ROUND(P{r}*(1+{P_RANGO}),2)",
            19: round(l["pmax"], 2), 20: None, 21: f'=IFERROR(IF({last}=0,"",{last}),"")',
            22: (f'=IF(AND(G{r}>={P_STOP},K{r}=0),"PAUSAR (≥15 clics sin venta)",'
                 f'IF(AND(E{r}<=Q{r},G{r}>=20,M{r}>0.5),"PAUSAR (puja mínima y ACOS>50%)",""))'),
            23: (f'=IF(LEFT(V{r},6)="PAUSAR","Pausar y crear ticket",IF(G{r}<{P_MIN},"Esperar (menos de 10 clics)",'
                 f'IF(K{r}=0,"Sin ventas aún: vigilar",IF(M{r}<0.3,"Subir 10-25% (máx. "&TEXT(MIN(R{r},S{r}),"0.00")&" €)",'
                 f'IF(M{r}<={P_ACOS},"En objetivo: no tocar",IF(M{r}<=0.5,"Bajar 10-25% (mín. "&TEXT(Q{r},"0.00")&" €)",'
                 f'"Bajar 25-50% (mín. "&TEXT(Q{r},"0.00")&" €)"))))))'),
            24: f'=A{r}&" | "&C{r}&" | "&D{r}',
        }
        for col, v in vals.items():
            put(wsg, r, col, v)
        for col in (5, 9, 10, 12, 16, 17, 18, 19):
            wsg.cell(row=r, column=col).number_format = EUR
        for col in (6, 7, 11):
            wsg.cell(row=r, column=col).number_format = INT
        for col in (8, 13):
            wsg.cell(row=r, column=col).number_format = PCT
        wsg.cell(row=r, column=14).number_format = X2
        wsg.cell(row=r, column=21).number_format = DATE
        for col in (3, 15):
            wsg.cell(row=r, column=col).fill = FILL_ORIGEN[l["origen"]]
        wsg.cell(row=r, column=16).font, wsg.cell(row=r, column=16).fill = f_input, fill_input
        l["fila"] = r
    tr = SL + 1
    for c in CAMPANAS:
        rows = [l["fila"] for l in lineas if l["campana"] == c["nombre"]]
        put(wsg, tr, 1, "TOTAL", font=f_bold)
        put(wsg, tr, 3, c["nombre"], font=f_bold)
        for col, L in ((6, "F"), (7, "G"), (9, "I"), (11, "K"), (12, "L")):
            put(wsg, tr, col, f"=SUM({L}{rows[0]}:{L}{rows[-1]})", EUR if L in "IL" else INT, f_bold)
        for col, fml, fmt in ((8, f"=IFERROR(G{tr}/F{tr},0)", PCT), (10, f"=IFERROR(I{tr}/G{tr},0)", EUR),
                              (13, f"=IFERROR(I{tr}/L{tr},0)", PCT), (14, f"=IFERROR(L{tr}/I{tr},0)", X2)):
            put(wsg, tr, col, fml, fmt, f_bold)
        tr += 1
    wsg.conditional_formatting.add(f"V{S0}:V{SL}", FormulaRule(formula=[f'LEFT(V{S0},6)="PAUSAR"'], fill=red))
    wsg.conditional_formatting.add(f"M{S0}:M{SL}", FormulaRule(formula=[f"AND(L{S0}>0,M{S0}<={P_ACOS})"], fill=green))
    wsg.conditional_formatting.add(f"M{S0}:M{SL}", FormulaRule(formula=[f"AND(L{S0}>0,M{S0}>{P_ACOS})"], fill=red))
    wsg.freeze_panes = "D4"
    wsg.column_dimensions["X"].hidden = True

    # ---------------- Seguimiento
    wf = wb.create_sheet("Seguimiento")
    fcols = ["Fecha", "Campaña", "Palabra clave / segmentación", "Tipo de coincidencia", "Estado", "Ticket", "Puja (EUR)", "Impresiones",
             "Clics", "Coste total (EUR)", "Compras", "Ventas (EUR)", "CTR", "CPC (EUR)", "ACOS", "ROAS", "Nota", "Clave", "Clave+fecha"]
    widths(wf, [11, 22, 46, 12, 10, 8, 10, 11, 8, 11, 9, 11, 8, 9, 9, 8, 36, 0.1, 0.1])
    wf["A1"], wf["A1"].font = "Seguimiento — una fila por keyword/ASIN en cada revisión (cada 3 días). Datos ACUMULADOS desde la creación.", f_bold
    header(wf, 3, fcols)
    FR = 600
    for i in range(FR):
        r = 4 + i
        base = {}
        if i < len(lineas):
            l = lineas[i]
            base = {1: hoy, 2: l["campana"], 3: l["texto"], 4: l["match"], 5: "Pausado", 7: l["puja"],
                    8: 0, 9: 0, 10: 0, 11: 0, 12: 0, 17: "Día 0: creación (sin datos todavía)"}
        for col in range(1, 18):
            c = put(wf, r, col, base.get(col), font=f_input if col not in (13, 14, 15, 16) else f_base)
            if i >= len(lineas) and col not in (13, 14, 15, 16):
                c.fill = fill_input
        put(wf, r, 13, f'=IF(C{r}="","",IFERROR(I{r}/H{r},0))', PCT)
        put(wf, r, 14, f'=IF(C{r}="","",IFERROR(J{r}/I{r},0))', EUR)
        put(wf, r, 15, f'=IF(C{r}="","",IFERROR(J{r}/L{r},0))', PCT)
        put(wf, r, 16, f'=IF(C{r}="","",IFERROR(L{r}/J{r},0))', X2)
        put(wf, r, 18, f'=IF(C{r}="","",B{r}&" | "&C{r}&" | "&D{r})')
        put(wf, r, 19, f'=IF(C{r}="","",R{r}&"#"&A{r})')
        wf.cell(row=r, column=1).number_format = DATE
        for col in (7, 10, 12):
            wf.cell(row=r, column=col).number_format = EUR
        for col in (8, 9, 11):
            wf.cell(row=r, column=col).number_format = INT
    for dv, rng in ((DataValidation(type="list", formula1=f"=Campañas!$A$4:$A${3 + len(CAMPANAS)}", allow_blank=True), f"B4:B{3 + FR}"),
                    (DataValidation(type="list", formula1='"Amplia,Frase,Exacta,Producto,Categoría"', allow_blank=True), f"D4:D{3 + FR}"),
                    (DataValidation(type="list", formula1='"Activado,Pausado,Archivado"', allow_blank=True), f"E4:E{3 + FR}")):
        wf.add_data_validation(dv)
        dv.add(rng)
    wf.freeze_panes = "D4"
    wf.column_dimensions["R"].hidden = wf.column_dimensions["S"].hidden = True

    # ---------------- Tickets
    wt = wb.create_sheet("Tickets")
    tcols = ["Ticket", "Fecha creación", "Campaña", "Tipo de cambio", "Palabra clave / segmentación", "Tipo de coincidencia",
             "Valor antes", "Valor después", "Cambio %", "Motivo / hipótesis", "ACOS de referencia", "Qué es la referencia",
             "Estado del ticket", "Fecha aplicado", "Revisar a partir de", "Base: Clics", "Base: Coste", "Base: Compras", "Base: Ventas",
             "Después: Clics", "Después: Coste", "Después: Compras", "Después: Ventas", "Después: ACOS", "Veredicto", "Cómo se aplica", "Notas"]
    widths(wt, [8, 11, 22, 24, 44, 12, 10, 10, 9, 60, 11, 22, 16, 11, 11, 9, 9, 9, 9, 10, 10, 10, 10, 10, 28, 14, 26])
    wt["A1"], wt["A1"].font = "Tickets — uno por cambio. 'Base' = acumulado al aplicar el cambio; 'Después' = solo lo ocurrido desde entonces.", f_bold
    header(wt, 3, tcols)
    tk = []  # (campaña, tipo, texto, match, antes, después, motivo, ref, ref_txt, como)
    tk.append(("—", "Crear cartera", "", "", None, CARTERA_MENSUAL,
               f"Cartera 'FreshFinder - cartera' con límite MENSUAL recurrente de {CARTERA_MENSUAL:.0f} € y las 3 campañas dentro. Es el tope duro: Amazon para todo al llegar.",
               None, "", "A mano"))
    for c in CAMPANAS:
        tk.append((c["nombre"], "Crear campaña y grupo", "", "", None, PRESUPUESTO_DIARIO,
                   f"Campaña '{c['nombre']}' (manual, {c['tipo'].lower()}, {PRESUPUESTO_DIARIO:.0f} €/día, pujas dinámicas solo reducir) "
                   f"con el grupo '{c['grupo']}' y el anuncio {c['asin']} ({c['sku']}). Se crea en pausa.", None, "", "Hoja masiva"))
        for l in (l for l in lineas if l["campana"] == c["nombre"]):
            tipo = {"Histórica": "Añadir keyword (histórica)", "Especial IA": "Añadir keyword (especial IA)"}.get(l["origen"], "Añadir segmentación (prueba)")
            ref, ref_txt = (l["ref"], "ACOS en el histórico") if l["ref"] != "obj" else (f"={P_ACOS}", "ACOS objetivo")
            tk.append((c["nombre"], tipo, l["texto"], l["match"], l["antes"], l["puja"], l["motivo"], ref, ref_txt,
                       "A mano" if l.get("manual") else "Hoja masiva"))
            l["ticket_idx"] = len(tk)
        for term, motivo in c.get("negativas", []):
            tk.append((c["nombre"], "Añadir negativa", term, "Frase negativa", None, None, motivo, None, "", "Hoja masiva"))
        for a, motivo in c.get("negativos_producto", []):
            tk.append((c["nombre"], "Añadir negativa", f'asin="{a}"', "Producto negativo", None, None, motivo, None, "", "Hoja masiva"))
    TR = len(tk) + 60
    for i in range(TR):
        r = 4 + i
        row = tk[i] if i < len(tk) else None
        vals = {1: f"T-{i + 1:03d}"}
        if row:
            camp, tipo, texto, mt, antes, desp, motivo, ref, ref_txt, como = row
            vals.update({2: hoy, 3: camp, 4: tipo, 5: texto, 6: mt, 7: antes, 8: desp, 10: motivo, 11: ref, 12: ref_txt,
                         13: "Pendiente de aplicar", 16: 0, 17: 0, 18: 0, 19: 0, 26: como})
        for col in range(1, 28):
            put(wt, r, col, vals.get(col), al=Alignment(vertical="top", wrap_text=col in (5, 10, 25, 27)))
        for col in ((13, 14, 27) if row else (2, 3, 4, 5, 6, 7, 8, 10, 11, 12, 13, 14, 16, 17, 18, 19, 26, 27)):
            wt.cell(row=r, column=col).font, wt.cell(row=r, column=col).fill = f_input, fill_input
        wt.cell(row=r, column=9, value=f'=IF(OR(G{r}="",H{r}="",G{r}=0),"",H{r}/G{r}-1)')
        wt.cell(row=r, column=15, value=f'=IF(N{r}="","",N{r}+{P_EVAL})')
        m = f'MATCH(C{r}&" | "&E{r}&" | "&F{r},Segmentación!$X${S0}:$X${SL},0)'
        for col, sc, bc in ((20, "G", "P"), (21, "I", "Q"), (22, "K", "R"), (23, "L", "S")):
            wt.cell(row=r, column=col, value=f'=IF(OR(E{r}="",N{r}=""),"",IFERROR(INDEX(Segmentación!${sc}${S0}:${sc}${SL},{m})-{bc}{r},""))')
        wt.cell(row=r, column=24, value=f'=IF(OR(T{r}="",W{r}=""),"",IF(W{r}=0,"",U{r}/W{r}))')
        wt.cell(row=r, column=25, value=(
            f'=IF(D{r}="","",IF(OR(E{r}="",ISNUMBER(SEARCH("negativa",D{r}))),IF(N{r}="","Pendiente de aplicar","Aplicado (no se evalúa)"),'
            f'IF(N{r}="","Pendiente de aplicar",IF(T{r}="","Sin datos",IF(T{r}<{P_MIN},"Esperando datos ("&T{r}&" de 10 clics)",'
            f'IF(V{r}=0,"Empeora: "&T{r}&" clics sin venta",IF(TODAY()<O{r},"Esperando atribución (7 días)",'
            f'IF(X{r}<=K{r},"Mejora (ACOS "&TEXT(X{r},"0%")&")","Empeora (ACOS "&TEXT(X{r},"0%")&")"))))))))'))
        for col in (9, 11, 24):
            wt.cell(row=r, column=col).number_format = PCT
        for col in (7, 8, 17, 19, 21, 23):
            wt.cell(row=r, column=col).number_format = EUR
        for col in (16, 18, 20, 22):
            wt.cell(row=r, column=col).number_format = INT
        for col in (2, 14, 15):
            wt.cell(row=r, column=col).number_format = DATE
        if row:
            wt.row_dimensions[r].height = 48
    for l in lineas:
        wsg.cell(row=l["fila"], column=20, value=f"T-{l['ticket_idx']:03d}").font = f_base
    for dv, rng in ((DataValidation(type="list", formula1='"Pendiente de aplicar,Aplicado,Evaluado,Descartado"', allow_blank=True), f"M4:M{3 + TR}"),
                    (DataValidation(type="list", formula1='"Subir puja,Bajar puja,Pausar keyword,Reactivar keyword,Añadir keyword,Añadir negativa,Añadir segmentación (prueba),Cambiar ajuste por ubicación,Cambiar presupuesto,Otro"', allow_blank=True), f"D{4 + len(tk)}:D{3 + TR}")):
        wt.add_data_validation(dv)
        dv.add(rng)
    wt.conditional_formatting.add(f"Y4:Y{3 + TR}", FormulaRule(formula=['LEFT(Y4,6)="Mejora"'], fill=green))
    wt.conditional_formatting.add(f"Y4:Y{3 + TR}", FormulaRule(formula=['LEFT(Y4,7)="Empeora"'], fill=red))
    wt.freeze_panes = "F4"

    # ---------------- Histórico
    wh = wb.create_sheet("Histórico")
    hcols = ["Producto", "Campaña / grupo de anuncios", "Palabra clave", "Tipo de coincidencia", "Puja (EUR)", "Impresiones", "Clics", "CTR",
             "Coste total (EUR)", "CPC (EUR)", "Compras", "Ventas (EUR)", "ACOS", "ROAS", "¿En la campaña nueva?"]
    widths(wh, [9, 44, 44, 12, 10, 11, 8, 8, 11, 9, 9, 11, 9, 8, 14])
    wh["A1"], wh["A1"].font = f"Histórico por producto ({EXCEL_HIST}; solo grupos donde el producto tiene ≥ 90% del gasto)", f_bold
    header(wh, 3, hcols)
    elegidas = {(l["campana"].split(" - ")[0].lower(), k.normalizar(l["texto"]), l["match"]) for l in lineas if l["origen"] == "Histórica"}
    orden = sorted(historico, key=lambda x: (x[0], (x[2]["coste"] / x[2]["ventas"]) if x[2]["ventas"] else 99 + (0 if x[2]["clics"] else 1)))
    for i, (perfil, asin, f) in enumerate(orden):
        r = 4 + i
        sel = (perfil, k.normalizar(f["keyword"]), f["coincidencia"]) in elegidas
        vals = [perfil, f["archivo"], f["keyword"], f["coincidencia"], f["puja"], f["impresiones"], f["clics"], f"=IFERROR(G{r}/F{r},0)",
                f["coste"], f"=IFERROR(I{r}/G{r},0)", f["compras"], f["ventas"], f"=IFERROR(I{r}/L{r},0)", f"=IFERROR(L{r}/I{r},0)", "Sí" if sel else ""]
        for col, v in enumerate(vals, 1):
            c = put(wh, r, col, v)
            if sel and col in (3, 15):
                c.fill = FILL_ORIGEN["Histórica"]
        for col in (5, 9, 10, 12):
            wh.cell(row=r, column=col).number_format = EUR
        for col in (6, 7, 11):
            wh.cell(row=r, column=col).number_format = INT
        for col in (8, 13):
            wh.cell(row=r, column=col).number_format = PCT
        wh.cell(row=r, column=14).number_format = X2
    wh.freeze_panes = "D4"
    wh.auto_filter.ref = f"A3:O{3 + len(orden)}"
    wb.save(salida)
    return len(tk)


# ================================================================== bulk
def crear_bulk(lineas, plantilla, salida, hoy):
    wb = load_workbook(plantilla)
    ws = wb["Camp. de Sponsored Products"]
    cols = [c.value for c in ws[1]]
    ini = f"{hoy:%Y%m%d}"

    def row(**kv):
        base = {"Producto": "Sponsored Products", "Operación": "Crear"}
        base.update(kv)
        ws.append([base.get(c) for c in cols])

    for c in CAMPANAS:
        C, G = c["nombre"], c["grupo"]
        ids = {"ID de la campaña": C, "ID del grupo de anuncios": G}
        row(Entidad="Campaña", **{"ID de la campaña": C, "Nombre de la campaña": C, "Fecha de inicio": ini,
                                   "Tipo de segmentación": "Manual", "Estado": "Pausado", "Presupuesto diario": PRESUPUESTO_DIARIO,
                                   "Estrategia de pujas": "Pujas dinámicas: solo reducir"})
        row(Entidad="Grupo de anuncios", **ids, **{"Nombre del grupo de anuncios": G, "Estado": "Habilitado",
                                                  "Puja predeterminada del grupo de anuncios": c["puja_grupo"]})
        row(Entidad="Anuncio de producto", **ids, **{"SKU": c["sku"], "Estado": "Habilitado"})
        for l in (l for l in lineas if l["campana"] == C and not l.get("manual")):
            if l["match"] == "Producto":
                row(Entidad="Segmentación por productos", **ids, **{"Estado": "Habilitado", "Puja": l["puja"],
                                                                   "Fórmula de segmentación por productos": l["texto"]})
            else:
                row(Entidad="Palabra clave", **ids, **{"Estado": "Habilitado", "Puja": l["puja"],
                                                       "Texto de palabra clave": l["texto"], "Tipo de coincidencia": l["match"]})
        for term, _ in c.get("negativas", []):
            row(Entidad="Palabra clave negativa", **ids, **{"Estado": "Habilitado", "Texto de palabra clave": term,
                                                            "Tipo de coincidencia": "Frase negativa"})
        for a, _ in c.get("negativos_producto", []):
            row(Entidad="Segmentación por productos negativa", **ids, **{"Estado": "Habilitado",
                                                                        "Fórmula de segmentación por productos": f'asin="{a}"'})
    wb.save(salida)
    return ws.max_row - 1


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--fecha", default=date.today().isoformat())
    ap.add_argument("--plantilla", default="plantillas/AdvertisingBulksheetTemplate-seller.xlsx")
    args = ap.parse_args()
    hoy = date.fromisoformat(args.fecha)
    productos, lineas, historico = preparar()
    n_tk = crear_memoria(productos, lineas, historico, hoy, "resultados/memoria.xlsx")
    n_bulk = crear_bulk(lineas, args.plantilla, f"resultados/bulk_{hoy:%Y-%m-%d}.xlsx", hoy)
    print(f"resultados/memoria.xlsx: {len(lineas)} segmentaciones, {n_tk} tickets")
    print(f"resultados/bulk_{hoy:%Y-%m-%d}.xlsx: {n_bulk} filas")
    for c in CAMPANAS:
        print(f"  {c['nombre']}:")
        for l in (l for l in lineas if l["campana"] == c["nombre"]):
            print(f"    {l['origen']:<16} {l['texto']:<52} {l['match']:<9} {l['puja']:.2f} €")


if __name__ == "__main__":
    main()
