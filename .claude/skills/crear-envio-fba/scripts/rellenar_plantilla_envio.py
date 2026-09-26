"""Paso 1: rellena la plantilla de envío de "Send to Amazon" (ManifestFileUpload_Template_MPL.xlsx).

Uso:
    python rellenar_plantilla_envio.py PLANTILLA.xlsx SPEC.json SALIDA.xlsx
    python rellenar_plantilla_envio.py PLANTILLA.xlsx --leer

SPEC.json:
{
  "prep_owner_por_defecto": "Seller",        # opcional; si falta se deja el de la plantilla
  "labeling_owner_por_defecto": "Seller",    # opcional
  "lineas": [
    {"sku": "5E-I8NY-S191", "cantidad": 20},
    {"sku": "O8-5W7J-DSK1", "cantidad": 20, "prep_owner": "Amazon",
     "unidades_por_caja": 10, "numero_cajas": 2,
     "caja": {"largo": 30, "ancho": 20, "alto": 15, "peso": 2.1}}
  ]
}

Solo escribe en las celdas de datos. No toca cabeceras, instrucciones ni el resto de pestañas.
Las columnas de caja/caducidad/lote solo se rellenan si la plantilla las trae (depende de las
opciones que se eligieron al descargarla en Seller Central).
"""

import sys
import warnings
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from common import Errores, cargar_spec, normalizar_owner, validar_caja, validar_sku  # noqa: E402

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Prefijo en minúsculas del nombre de columna de Amazon -> clave interna
COLUMNAS = {
    "merchant sku": "sku",
    "quantity": "cantidad",
    "prep owner": "prep_owner",
    "labeling owner": "labeling_owner",
    "expiration date": "caducidad",
    "manufacturing lot code": "lote",
    "units per box": "unidades_por_caja",
    "number of boxes": "numero_cajas",
    "box length": "largo",
    "box width": "ancho",
    "box height": "alto",
    "box weight": "peso",
}


def hoja_plantilla(wb):
    for ws in wb.worksheets:
        nombre = ws.title.lower()
        if nombre.startswith("create workflow") and "template" in nombre:
            return ws
    sys.exit("No encuentro la pestaña 'Create workflow – template'. ¿Es la plantilla correcta?")


def localizar(ws):
    """Devuelve (fila_cabecera, {clave: columna}, {'prep': fila, 'labeling': fila})."""
    fila_cab, defaults = None, {}
    for row in ws.iter_rows(min_col=1, max_col=1):
        c = row[0]
        v = str(c.value or "").strip().lower()
        if v == "default prep owner":
            defaults["prep"] = c.row
        elif v == "default labeling owner":
            defaults["labeling"] = c.row
        elif v == "merchant sku":
            fila_cab = c.row
            break
    if fila_cab is None:
        sys.exit("No encuentro la fila de cabecera 'Merchant SKU'.")
    cols = {}
    for c in ws[fila_cab]:
        v = str(c.value or "").strip().lower()
        for prefijo, clave in COLUMNAS.items():
            if v.startswith(prefijo):
                cols[clave] = c.column
    return fila_cab, cols, defaults


def leer(ruta):
    ws = hoja_plantilla(openpyxl.load_workbook(ruta))
    fila_cab, cols, defaults = localizar(ws)
    print(f"Pestaña: {ws.title}")
    for k, fila in defaults.items():
        print(f"Default {k} owner (fila {fila}): {ws.cell(fila, 2).value}")
    print(f"Cabecera en fila {fila_cab}. Columnas disponibles: {', '.join(cols)}")
    faltan = [k for k in ("unidades_por_caja", "numero_cajas", "largo", "ancho", "alto", "peso") if k not in cols]
    if faltan:
        print("Esta plantilla NO lleva datos de caja: se darán en el paso 2 (contenido de cajas).")


def rellenar(ruta, spec, salida):
    wb = openpyxl.load_workbook(ruta)
    ws = hoja_plantilla(wb)
    fila_cab, cols, defaults = localizar(ws)
    err = Errores()

    for clave, fila_key in (("prep_owner_por_defecto", "prep"), ("labeling_owner_por_defecto", "labeling")):
        valor = normalizar_owner(spec.get(clave), clave, err)
        if valor is not None:
            if fila_key not in defaults:
                err.error(f"La plantilla no tiene la fila '{clave}'.")
            else:
                ws.cell(defaults[fila_key], 2, valor)

    lineas = spec.get("lineas") or []
    if not lineas:
        err.error("La spec no tiene 'lineas'.")
    usos = {}
    for i, ln in enumerate(lineas, start=1):
        sku = ln.get("sku")
        validar_sku(sku, err)
        usos[sku] = usos.get(sku, 0) + 1
        q = ln.get("cantidad")
        if not isinstance(q, int) or q < 1:
            err.error(f"Línea {i} ({sku}): la cantidad debe ser un entero >= 1 (es {q!r}).")
        for campo in ("prep_owner", "labeling_owner"):
            ln[campo] = normalizar_owner(ln.get(campo), f"Línea {i} {campo}", err)
        upb, nb = ln.get("unidades_por_caja"), ln.get("numero_cajas")
        if upb is not None or nb is not None:
            falta_col = [k for k in ("unidades_por_caja", "numero_cajas", "largo", "ancho", "alto", "peso") if k not in cols]
            if falta_col:
                err.error(f"Línea {i} ({sku}) trae datos de caja pero la plantilla no tiene esas columnas. "
                          "Descarga la plantilla con información de embalaje o dalos en el paso 2.")
            if not (isinstance(upb, int) and isinstance(nb, int) and upb >= 1 and nb >= 1):
                err.error(f"Línea {i} ({sku}): unidades_por_caja y numero_cajas deben ser enteros >= 1.")
            elif isinstance(q, int) and upb * nb != q:
                err.error(f"Línea {i} ({sku}): {upb} uds/caja × {nb} cajas = {upb * nb}, no cuadra con cantidad {q}.")
            caja = ln.get("caja") or {}
            validar_caja(f"Línea {i} ({sku})", caja.get("peso"), caja.get("largo"), caja.get("ancho"),
                         caja.get("alto"), err)
        for campo in ("caducidad", "lote"):
            if ln.get(campo) is not None and campo not in cols:
                err.error(f"Línea {i} ({sku}): la plantilla no tiene columna para '{campo}'.")
    for sku, n in usos.items():
        if n > 4:
            err.error(f"{sku} aparece {n} veces; Amazon admite como máximo 4 filas por SKU.")

    # Las filas de datos tienen que estar vacías: no pisamos nada que ya esté escrito.
    for r in range(fila_cab + 1, fila_cab + 1 + len(lineas)):
        if any(ws.cell(r, c).value not in (None, "") for c in cols.values()):
            err.error(f"La fila {r} ya tiene datos. Usa una plantilla limpia.")

    err.informar_y_salir_si_hay_errores()

    for r, ln in enumerate(lineas, start=fila_cab + 1):
        ws.cell(r, cols["sku"], ln["sku"])
        ws.cell(r, cols["cantidad"], ln["cantidad"])
        for campo in ("prep_owner", "labeling_owner", "caducidad", "lote", "unidades_por_caja", "numero_cajas"):
            if ln.get(campo) is not None and campo in cols:
                ws.cell(r, cols[campo], ln[campo])
        for campo, valor in (ln.get("caja") or {}).items():
            if campo in cols:
                ws.cell(r, cols[campo], valor)

    Path(salida).parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    total = sum(ln["cantidad"] for ln in lineas)
    print(f"OK: {salida} — {len(lineas)} línea(s), {total} unidades.")


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[2] == "--leer":
        leer(sys.argv[1])
    elif len(sys.argv) == 4:
        rellenar(sys.argv[1], cargar_spec(sys.argv[2]), sys.argv[3])
    else:
        sys.exit(__doc__)
