"""
fuente_bulk.py — fuente alternativa SIN API: la hoja masiva descargada de Amazon Ads.

Mientras no estén las credenciales de la API (o si un día falla), el agente puede trabajar
con la descarga de Operaciones en bloque (Sponsored Products, rango desde la creación de las
campañas hasta hoy, con campañas en pausa) que se deja en entradas/AAAA-MM-DD/:

  LEER    -> la misma estructura Cuenta que da ads_api.py, con los acumulados de cada fila
             (se guardan como "foto" del día en la hoja Seguimiento).
  APLICAR -> no toca Amazon: escribe salidas/AAAA-MM-DD/bulk_cambios_<fecha>.xlsx sobre la plantilla
             oficial, para subirlo a mano. Los tickets quedan "enviado_bulk" y se CONFIRMAN
             (o se marcan fallidos) en la siguiente ejecución, al leer la nueva descarga (§2.8).

Las cabeceras se buscan por nombre (sin mayúsculas ni acentos), porque la descarga trae
columnas "(Solo informativo)" que cambian con el tiempo. Los valores que se escriben (Producto,
Estado) copian los de la propia descarga: lo que Amazon da es lo que Amazon acepta de vuelta.
Operación: Crear / Actualizar / Archivar (vacía = Amazon ignora la fila; Amazon Ads Academy).

Además de las campañas se leen las palabras clave negativas, las métricas de cada emplazamiento
(filas "Ajuste de puja") y la hoja "Inf. de Térm. de Búsq. de SP" (términos de búsqueda: para
descargarla hay que marcar ese informe al pedir la hoja masiva).
"""

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

import config
import keyword_ml as kml
import pujas
from modelo import (ACTIVO, ARCHIVADO, AUTO, CATEGORIA, EMPLAZAMIENTO, ESTRATEGIA, FINALIZADA, KEYWORD, NEGATIVA,
                    NUEVA_KEYWORD, NUEVO_ASIN, PAUSADO, PAUSAR, PRESUPUESTO, PRODUCTO, PUJA, REACTIVAR, REACTIVAR_CAMPANA,
                    SOLO_BAJA, CREAR_CAMPANA, Anuncio, Campana, Cuenta, Elemento, Grupo, Metricas, Negativa, Termino)

HOJA = "Camp. de Sponsored Products"
HOJA_TERMINOS = "Inf. de Térm. de Búsq. de SP"
PLANTILLA = config.RAIZ / "plantillas" / "AdvertisingBulksheetTemplate-seller.xlsx"
OP_ACTUALIZAR = "Actualizar"   # confirmado: Crear / Actualizar / Archivar (Amazon Ads Academy)
ACTIVADO, EN_PAUSA = "activada", "en pausa"    # como vienen en la descarga de Amazon.es


def _estado(v):
    s = kml.normalizar(str(v or ""))
    if "archiv" in s:
        return ARCHIVADO
    if "pausa" in s or "paused" in s:
        return PAUSADO
    if "activ" in s or "enabled" in s or "habilit" in s:
        return ACTIVO
    return PAUSADO


def _fecha(v):
    """'20250121', una fecha de Excel o nada."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.fullmatch(r"(20\d\d)-?(\d\d)-?(\d\d)", str(v or "").strip())
    return date(int(m[1]), int(m[2]), int(m[3])) if m else None


def _num(v):
    if v in (None, "", "-"):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    return kml._num(str(v).replace("%", ""))


class _Cols:
    def __init__(self, cab):
        self.cab = [kml.normalizar(str(c or "")) for c in cab]

    def i(self, *nombres, exacto=False):
        for n in nombres:
            n = kml.normalizar(n)
            for j, c in enumerate(self.cab):
                if (c == n) if exacto else (c.startswith(n)):
                    return j
        return None


def buscar_descarga(carpeta):
    """La hoja masiva más reciente de una carpeta (xlsx con la hoja de Sponsored Products)."""
    carpeta = Path(carpeta)
    for p in sorted(carpeta.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True):
        try:
            if HOJA in load_workbook(p, read_only=True).sheetnames:
                return p
        except Exception:
            continue
    return None


def fecha_de_archivo(ruta):
    m = re.search(r"(20\d\d)[-_]?(\d\d)[-_]?(\d\d)", Path(ruta).name)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    return date.fromtimestamp(Path(ruta).stat().st_mtime)


class FuenteBulk:
    nombre = "hoja masiva"
    verifica_al_momento = False

    def __init__(self, ruta, salida_dir):
        self.ruta = Path(ruta)
        self.salida_dir = Path(salida_dir)
        self.pendientes = []
        self.producto_txt, self.activada, self.en_pausa = "Sponsored Products", ACTIVADO, EN_PAUSA

    def leer_cuenta(self, hoy=None):
        # sin read_only: algunas descargas de Amazon traen mal la etiqueta interna de
        # dimensiones de la hoja (dice A1:A1 aunque hay cientos de filas reales), y en modo
        # read_only openpyxl confía en esa etiqueta y se queda solo con la primera celda.
        wb = load_workbook(self.ruta, data_only=True)
        filas = list(wb[HOJA].iter_rows(values_only=True))
        k = _Cols(filas[0])
        col = {
            "entidad": k.i("entidad", exacto=True), "id_c": k.i("id de la campana"), "id_g": k.i("id del grupo de anuncios"),
            "id_cart": k.i("id de la cartera"), "id_a": k.i("id del anuncio"), "id_k": k.i("id de palabra clave"),
            "id_t": k.i("id de segmentacion por productos"), "nombre_c": k.i("nombre de la campana"),
            "nombre_g": k.i("nombre del grupo de anuncios"), "tipo_seg": k.i("tipo de segmentacion"),
            "estado": k.i("estado", exacto=True), "estado_serv": k.i("estado de publicacion", "estado de entrega",
                                                                   "estado de elegibilidad"),
            "presupuesto": k.i("presupuesto diario", "presupuesto"), "sku": k.i("sku"), "asin": k.i("asin"),
            "puja_g": k.i("puja predeterminada"), "puja": k.i("puja", exacto=True), "texto": k.i("texto de palabra clave"),
            "coinc": k.i("tipo de coincidencia"), "formula": k.i("formula de segmentacion por productos"),
            "estrategia": k.i("estrategia de pujas"), "emplazamiento": k.i("emplazamiento", exacto=True),
            "porcentaje": k.i("porcentaje", exacto=True), "fin": k.i("fecha de finalizacion"),
            "resuelta": k.i("texto de expresion resuelta", "expresion resuelta"),
            "impr": k.i("impresiones"), "clics": k.i("clics"), "gasto": k.i("gasto", "coste", "inversion"),
            "ventas": k.i("ventas"), "pedidos": k.i("pedidos", "compras"),
        }
        v = lambda r, c: r[col[c]] if col[c] is not None and col[c] < len(r) else None
        s = lambda r, c: str(v(r, c)).strip() if v(r, c) not in (None, "") else ""
        cuenta = Cuenta(fecha=hoy or date.today())
        ajustes = []
        prod_col = k.i("producto", exacto=True)
        for r in filas[1:]:
            if prod_col is not None and r[prod_col]:
                self.producto_txt = str(r[prod_col])
                break
        for r in filas[1:]:
            est = s(r, "estado")
            if est and _estado(est) == ACTIVO:
                self.activada = est
            elif "pausa" in est.lower():
                self.en_pausa = est
        for r in filas[1:]:
            ent = kml.normalizar(s(r, "entidad"))
            met = Metricas(_num(v(r, "clics")), _num(v(r, "gasto")), _num(v(r, "pedidos")), _num(v(r, "ventas")),
                           _num(v(r, "impr")))
            if ent == "campana":
                cuenta.campanas[s(r, "id_c")] = Campana(
                    id=s(r, "id_c"), nombre=s(r, "nombre_c") or s(r, "id_c"), estado=_estado(v(r, "estado")),
                    presupuesto=_num(v(r, "presupuesto")),
                    segmentacion="AUTO" if "autom" in kml.normalizar(s(r, "tipo_seg")) else "MANUAL",
                    id_cartera=s(r, "id_cart") or None, estrategia_pujas=pujas.normalizar_estrategia(v(r, "estrategia")),
                    fecha_fin=_fecha(v(r, "fin")))
            elif ent == "ajuste de puja":
                ajustes.append((s(r, "id_c"), pujas.normalizar_emplazamiento(v(r, "emplazamiento")), _num(v(r, "porcentaje")),
                                met))
            elif ent.startswith("palabra clave negativa"):
                cuenta.negativas.append(Negativa(clave=s(r, "id_k"), id_campana=s(r, "id_c"), id_grupo=s(r, "id_g"),
                                                 texto=s(r, "texto"), coincidencia=s(r, "coinc"),
                                                 estado=_estado(v(r, "estado"))))
            elif ent == "grupo de anuncios":
                cuenta.grupos[s(r, "id_g")] = Grupo(id=s(r, "id_g"), id_campana=s(r, "id_c"),
                                                    nombre=s(r, "nombre_g") or s(r, "id_g"), estado=_estado(v(r, "estado")),
                                                    puja_defecto=_num(v(r, "puja_g")) or None)
            elif ent == "anuncio de producto":
                cuenta.anuncios.append(Anuncio(id=s(r, "id_a"), id_campana=s(r, "id_c"), id_grupo=s(r, "id_g"),
                                               asin=s(r, "asin"), sku=s(r, "sku"), estado=_estado(v(r, "estado")),
                                               estado_servicio=s(r, "estado_serv")))
            elif ent == "palabra clave":
                cuenta.elementos[s(r, "id_k")] = Elemento(
                    clave=s(r, "id_k"), tipo=KEYWORD, id_campana=s(r, "id_c"), id_grupo=s(r, "id_g"),
                    texto=s(r, "texto"), coincidencia=s(r, "coinc").capitalize(), estado=_estado(v(r, "estado")),
                    puja=_num(v(r, "puja")) or None, metricas=met)
            elif ent == "segmentacion por productos":
                formula = s(r, "formula") or s(r, "resuelta")
                f = formula.lower()
                if f.startswith("asin"):
                    tipo, texto, coinc = PRODUCTO, formula.split("=", 1)[-1].strip('" ').upper(), "ASIN"
                elif f.startswith("category"):
                    tipo, texto, coinc = CATEGORIA, formula, "Categoría"
                else:
                    tipo, texto, coinc = AUTO, formula or "automática", "Automática"
                cuenta.elementos[s(r, "id_t")] = Elemento(
                    clave=s(r, "id_t"), tipo=tipo, id_campana=s(r, "id_c"), id_grupo=s(r, "id_g"), texto=texto,
                    coincidencia=coinc, estado=_estado(v(r, "estado")), puja=_num(v(r, "puja")) or None, metricas=met)
        for id_c, lugar, pct, met in ajustes:
            if id_c in cuenta.campanas and lugar:
                cuenta.campanas[id_c].ajustes_emplazamiento[lugar] = pct
                cuenta.campanas[id_c].metricas_emplazamiento[lugar] = met
        for c in cuenta.campanas.values():
            if c.estado == ACTIVO and c.fecha_fin and c.fecha_fin < cuenta.fecha:
                c.estado = FINALIZADA
        # la puja por defecto del grupo cuando la fila no trae puja propia
        for e in cuenta.elementos.values():
            if e.puja is None and e.id_grupo in cuenta.grupos:
                e.puja = cuenta.grupos[e.id_grupo].puja_defecto
        if HOJA_TERMINOS in wb.sheetnames:
            cuenta.terminos = leer_terminos(wb[HOJA_TERMINOS], cuenta)
        return cuenta

    def acumulados(self, cuenta, series=None, hoy=None):
        return {k: e.metricas for k, e in cuenta.elementos.items()}

    # ------------------------------------------------------------ "aplicar" = preparar hoja masiva
    def aplicar(self, cambio, cuenta):
        self.pendientes.append((cambio, cuenta))
        cambio.estado = "enviado_bulk"
        cambio.detalle = "En la hoja masiva de cambios: se confirmará al leer la próxima descarga"
        return cambio

    def cerrar(self, hoy):
        if not self.pendientes:
            return None
        wb = load_workbook(PLANTILLA)
        ws = wb[HOJA]
        cols = [c.value for c in ws[1]]

        activada, en_pausa = self.activada, self.en_pausa

        def fila(**kv):
            base = {"Producto": self.producto_txt} | kv
            ws.append([base.get(c) for c in cols])

        for c, cuenta in self.pendientes:
            ids = {"ID de la campaña": c.id_campana, "ID del grupo de anuncios": c.id_grupo}
            el = cuenta.elementos.get(c.clave)
            if c.tipo in (PUJA, PAUSAR, REACTIVAR) and el is not None:
                cambio = ({"Puja": c.despues} if c.tipo == PUJA else {"Estado": en_pausa} if c.tipo == PAUSAR
                          else {"Estado": activada, "Puja": c.despues})
                if el.tipo == KEYWORD:
                    fila(Entidad="Palabra clave", Operación=OP_ACTUALIZAR, **ids, **{"ID de palabra clave": el.clave}, **cambio)
                else:
                    fila(Entidad="Segmentación por productos", Operación=OP_ACTUALIZAR, **ids,
                         **{"ID de segmentación por productos": el.clave}, **cambio)
            elif c.tipo == PRESUPUESTO:
                fila(Entidad="Campaña", Operación=OP_ACTUALIZAR, **{"ID de la campaña": c.id_campana,
                                                                     "Presupuesto diario": c.despues})
            elif c.tipo == ESTRATEGIA:
                fila(Entidad="Campaña", Operación=OP_ACTUALIZAR, **{"ID de la campaña": c.id_campana,
                                                                     "Estrategia de pujas": c.despues})
            elif c.tipo == REACTIVAR_CAMPANA:
                fila(Entidad="Campaña", Operación=OP_ACTUALIZAR, **{"ID de la campaña": c.id_campana, "Estado": activada,
                                                                     "Presupuesto diario": c.extra.get("presupuesto")})
                for id_g in c.extra.get("grupos", []):
                    fila(Entidad="Grupo de anuncios", Operación=OP_ACTUALIZAR, **{
                        "ID de la campaña": c.id_campana, "ID del grupo de anuncios": id_g, "Estado": activada})
                for a in cuenta.anuncios:
                    if a.id in c.extra.get("anuncios", []):
                        fila(Entidad="Anuncio de producto", Operación=OP_ACTUALIZAR, **{
                            "ID de la campaña": c.id_campana, "ID del grupo de anuncios": a.id_grupo,
                            "ID del anuncio": a.id, "Estado": activada})
            elif c.tipo == NEGATIVA:
                fila(Entidad="Palabra clave negativa", Operación="Crear", **ids, **{
                    "Estado": activada, "Texto de palabra clave": c.texto, "Tipo de coincidencia": c.coincidencia})
            elif c.tipo == NUEVA_KEYWORD:
                fila(Entidad="Palabra clave", Operación="Crear", **ids, **{
                    "Estado": activada, "Puja": c.despues, "Texto de palabra clave": c.texto,
                    "Tipo de coincidencia": c.coincidencia})
            elif c.tipo == NUEVO_ASIN:
                fila(Entidad="Segmentación por productos", Operación="Crear", **ids, **{
                    "Estado": activada, "Puja": c.despues, "Fórmula de segmentación por productos": f'asin="{c.texto}"'})
            elif c.tipo == CREAR_CAMPANA:
                x = c.extra
                ref = {"ID de la campaña": x["nombre"], "ID del grupo de anuncios": x["nombre_grupo"]}
                fila(Entidad="Campaña", Operación="Crear", **{
                    "ID de la campaña": x["nombre"], "Nombre de la campaña": x["nombre"], "Fecha de inicio": f"{hoy:%Y%m%d}",
                    "Tipo de segmentación": "Manual", "Estado": activada, "Presupuesto diario": x["presupuesto"],
                    "Estrategia de pujas": SOLO_BAJA, "ID de la cartera": x.get("id_cartera")})
                fila(Entidad="Grupo de anuncios", Operación="Crear", **ref, **{
                    "Nombre del grupo de anuncios": x["nombre_grupo"], "Estado": activada,
                    "Puja predeterminada del grupo de anuncios": x["puja_grupo"]})
                fila(Entidad="Anuncio de producto", Operación="Crear", **ref, **{"SKU": x["sku"], "Estado": activada})
                for kw in x["keywords"]:
                    fila(Entidad="Palabra clave", Operación="Crear", **ref, **{
                        "Estado": activada, "Puja": kw["puja"], "Texto de palabra clave": kw["texto"],
                        "Tipo de coincidencia": kw["coincidencia"]})
        self.salida_dir.mkdir(parents=True, exist_ok=True)
        salida = self.salida_dir / f"bulk_cambios_{hoy:%Y-%m-%d}.xlsx"
        n = 1
        while salida.exists():
            n += 1
            salida = self.salida_dir / f"bulk_cambios_{hoy:%Y-%m-%d}_{n}.xlsx"
        wb.save(salida)
        return salida


def leer_terminos(ws, cuenta):
    """Filas de la hoja de términos de búsqueda -> [Termino] (acumulados del rango de la descarga)."""
    filas = list(ws.iter_rows(values_only=True))
    if len(filas) < 2:
        return []
    k = _Cols(filas[0])
    col = {"id_c": k.i("id de la campana"), "id_g": k.i("id del grupo de anuncios"), "id_k": k.i("id de palabra clave"),
           "id_t": k.i("id de segmentacion por productos"), "texto": k.i("texto de palabra clave"),
           "coinc": k.i("tipo de coincidencia"), "formula": k.i("formula de segmentacion por productos"),
           "termino": k.i("termino de busqueda"), "impr": k.i("impresiones"), "clics": k.i("clics"),
           "gasto": k.i("inversion", "gasto", "coste"), "ventas": k.i("ventas"), "pedidos": k.i("pedidos", "compras")}
    v = lambda r, c: r[col[c]] if col[c] is not None and col[c] < len(r) else None
    s = lambda r, c: str(v(r, c)).strip() if v(r, c) not in (None, "") else ""
    out = {}
    for r in filas[1:]:
        termino = s(r, "termino")
        if not termino or not s(r, "id_g"):
            continue
        clave_origen = s(r, "id_k") or s(r, "id_t")
        el = cuenta.elementos.get(clave_origen)
        origen = s(r, "texto") or s(r, "formula") or (el.texto if el else "")
        coinc = (s(r, "coinc").capitalize() if s(r, "coinc") else (el.coincidencia if el else "Automática"))
        met = Metricas(_num(v(r, "clics")), _num(v(r, "gasto")), _num(v(r, "pedidos")), _num(v(r, "ventas")),
                       _num(v(r, "impr")))
        t = Termino(id_campana=s(r, "id_c"), id_grupo=s(r, "id_g"), clave_origen=clave_origen, origen=origen,
                    coincidencia=coinc, termino=termino, metricas=met)
        if t.clave in out:                  # el mismo término cazado por varias keywords del grupo: se suma
            out[t.clave].metricas = out[t.clave].metricas + met
        else:
            out[t.clave] = t
    return list(out.values())
