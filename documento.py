"""
documento.py — el documento único del agente (AGENTE_AUTONOMO.md §4).

Un solo Excel con todo lo que el agente sabe y ha hecho:

  Leyenda        qué es cada hoja (para Juan)
  Resumen        la última ronda en cifras (se reescribe en cada ejecución)
  Campañas       estado, fondo (probado / experimentación), presupuesto y puntuación de cada campaña
  Segmentación   estado actual de cada keyword / ASIN / categoría, como los exports de Amazon
  Seguimiento    una FOTO por día y elemento: acumulados (clics, coste, compras, ventas)
  Diario         datos por día que da la API de Amazon (informe diario); si existen, mandan sobre las fotos
  Tickets        un registro por cambio: antes -> después, motivo, estado de verificación y veredicto
  Competencia    ASIN de la competencia por producto. Los edita JUAN a mano (§2.10)
  Investigación  frases candidatas que propone el módulo de investigación (§2.11)
  Alertas        todo lo que se ha avisado por correo (o quedó pendiente de enviar)
  Histórico      el rendimiento de antes del agente (FreshFinder_Amazon_Ads_historico.xlsx)

Escritura atómica: se escribe a un archivo temporal en la misma carpeta y solo cuando está
completo se sustituye el documento real (os.replace es atómico). Antes se guarda una copia
.bak de la versión anterior. Un corte a mitad nunca deja el documento roto.
"""

import json
import os
import tempfile
from bisect import bisect_right
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from modelo import Metricas

COLUMNAS = {
    "Resumen": ["Concepto", "Valor"],
    "Campañas": ["ID campaña", "Campaña", "Estado", "Producto (ASIN)", "Fondo", "Presupuesto diario (€)",
                 "Presupuesto objetivo (€)", "ACOS 30 días", "Puntuación", "Confianza (aprendizaje)",
                 "Creada por el agente", "Actualizado"],
    "Segmentación": ["Clave", "ID campaña", "Campaña", "ID grupo", "Grupo", "Producto (ASIN)", "Tipo",
                     "Palabra clave / segmentación", "Coincidencia", "Estado", "Puja (€)", "Elegible",
                     "Clics", "Coste (€)", "Compras", "Ventas (€)", "ACOS", "Clics maduros",
                     "Último cambio", "Clics desde el cambio", "Decisión de esta ronda",
                     "Requiere revisión de Juan", "Actualizado"],
    "Seguimiento": ["Fecha", "Clave", "ID campaña", "Producto (ASIN)", "Campaña", "Palabra clave / segmentación",
                    "Coincidencia", "Estado", "Puja (€)", "Impresiones", "Clics", "Coste (€)", "Compras",
                    "Ventas (€)"],
    "Diario": ["Fecha", "Clave", "ID campaña", "Impresiones", "Clics", "Coste (€)", "Compras", "Ventas (€)"],
    "Tickets": ["Ticket", "Fecha", "Tipo", "Clave", "Producto (ASIN)", "ID campaña", "Campaña", "ID grupo",
                "Palabra clave / segmentación", "Coincidencia", "Antes", "Después", "Motivo", "Estado",
                "Detalle", "Requiere revisión", "Base: clics", "Base: coste", "Base: compras", "Base: ventas",
                "Veredicto", "Fecha veredicto", "Después: clics", "Después: coste", "Después: compras",
                "Después: ventas", "ACOS después", "Datos extra"],
    "Competencia": ["Producto (ASIN)", "ASIN competidor", "Motivo", "Confirmado por Juan", "Añadido"],
    "Investigación": ["Fecha", "Producto (ASIN)", "Rank", "Palabra clave", "Motivo", "Fuente", "Volumen"],
    "Alertas": ["Fecha", "Hora", "Tipo", "Clave", "Mensaje", "Correo"],
    "Histórico": ["Producto (ASIN)", "Campaña", "Grupo de anuncios", "Keyword / segmento", "Coincidencia",
                  "Puja actual (€)", "Puja rec. baja", "Puja rec. mediana", "Impresiones", "Clics", "Coste (€)",
                  "Compras", "Ventas (€)"],
}

LEYENDA = [
    ("Documento único del agente autónomo de Amazon Ads (FreshFinder).", ""),
    ("Lo escribe agente.py en cada ejecución. Solo edita a mano: 'Competencia' (y, si quieres, la columna", ""),
    ("'Requiere revisión de Juan' la puedes leer para saber qué mirar). Cierra el Excel mientras corre el agente.", ""),
    ("", ""),
    ("Resumen", "La última ronda en cifras."),
    ("Campañas", "Fondo 'probado' (80 % del presupuesto) o 'experimentación' (20 %), presupuesto y puntuación."),
    ("Segmentación", "Estado actual de cada keyword / ASIN y la decisión de esta ronda."),
    ("Seguimiento", "Una foto por día: acumulados. 'Clics nuevos' = foto de hoy − foto del último cambio."),
    ("Diario", "Datos por día de la API. Permite saber qué clics tienen más de 7 días (maduros)."),
    ("Tickets", "Cada cambio: antes -> después, motivo, si se CONFIRMÓ releyendo Amazon, y el veredicto."),
    ("Competencia", "ASIN de la competencia. Pon 'Sí' en 'Confirmado por Juan' para que el agente pueda usarlos."),
    ("Investigación", "Frases candidatas del módulo de investigación (LLM). Solo son datos: las reglas deciden."),
    ("Alertas", "Avisos enviados (o pendientes de enviar si no hay correo configurado)."),
    ("Histórico", "Rendimiento de antes del agente, del Excel histórico."),
    ("", ""),
    ("Reglas", "Ronda: ≥3 días y ≥10 clics nuevos desde el último cambio. Clic maduro: >7 días."),
    ("", "Puja = P(compra|clic) × ticket medio × ACOS objetivo (30-35 %). Suelo 0,02 €. Solo reducir."),
    ("", "Stop-loss: 20 clics maduros o 4 € maduros sin ninguna venta -> pausa (nunca borrar)."),
    ("", "Máx. 12 keywords por grupo. Keyword nueva solo con ACOS predicho ≤ 30 %."),
    ("", "Tope 840 €/mes. 20 % experimentación, 80 % a lo probado según puntuación."),
]

FECHAS = {"Fecha", "Fecha veredicto", "Añadido", "Actualizado", "Último cambio"}


def fecha(v):
    """Acepta date, datetime o texto ISO y devuelve date (o None)."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v)[:10])


def num(v, defecto=0.0):
    if v is None or v == "":
        return defecto
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return defecto


class Documento:
    def __init__(self, ruta=None):
        self.ruta = Path(ruta or config.DOCUMENTO)
        self.hojas = {h: [] for h in COLUMNAS}
        self.nuevo = not self.ruta.exists()
        if not self.nuevo:
            self._cargar()

    # ------------------------------------------------------------ lectura / escritura
    def _cargar(self):
        wb = load_workbook(self.ruta, data_only=True)
        for h, cols in COLUMNAS.items():
            if h not in wb.sheetnames:
                continue
            it = wb[h].iter_rows(values_only=True)
            cab = [str(c) if c is not None else "" for c in next(it, [])]
            filas = []
            for r in it:
                if not any(v not in (None, "") for v in r):
                    continue
                d = dict(zip(cab, r))
                filas.append({c: d.get(c) for c in cols} | {k: v for k, v in d.items() if k not in cols and k})
            self.hojas[h] = filas

    def guardar(self):
        wb = Workbook()
        wb.remove(wb.active)
        cab_font, cab_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
        ws = wb.create_sheet("Leyenda")
        for a, b in LEYENDA:
            ws.append([a, b])
        ws["A1"].font = Font(bold=True, size=13)
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 110
        for h, cols in COLUMNAS.items():
            ws = wb.create_sheet(h)
            extra = []  # columnas que Juan haya añadido a mano: se conservan al final
            for f in self.hojas[h]:
                extra += [k for k in f if k not in cols and k not in extra]
            todas = cols + extra
            ws.append(todas)
            for c in ws[1]:
                c.font, c.fill = cab_font, cab_fill
            for f in self.hojas[h]:
                ws.append([_celda(f.get(c)) for c in todas])
            ws.freeze_panes = "A2"
            for i, c in enumerate(todas, 1):
                ws.column_dimensions[get_column_letter(i)].width = max(10, min(45, len(c) + 4))
        self.ruta.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(prefix=".tmp_", suffix=".xlsx", dir=self.ruta.parent)
        os.close(fd)
        try:
            wb.save(tmp)
            load_workbook(tmp, read_only=True).close()   # comprobar que se puede volver a abrir
            if self.ruta.exists():
                bak = self.ruta.with_suffix(".bak.xlsx")
                with open(self.ruta, "rb") as src, open(bak, "wb") as dst:
                    dst.write(src.read())
            os.replace(tmp, self.ruta)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
        self.nuevo = False

    # ------------------------------------------------------------ siembra (primera vez)
    def sembrar(self, hoy):
        """Primera ejecución: copia el histórico y la competencia de las campañas ya diseñadas."""
        if not self.hojas["Histórico"] and config.HISTORICO_XLSX.exists():
            wb = load_workbook(config.HISTORICO_XLSX, read_only=True, data_only=True)
            if "Keywords" in wb.sheetnames:
                it = wb["Keywords"].iter_rows(values_only=True)
                cab = [str(c or "").strip() for c in next(it)]
                for r in it:
                    d = dict(zip(cab, r))
                    if not d.get("Keyword / segmento"):
                        continue
                    self.hojas["Histórico"].append({c: d.get(c) for c in COLUMNAS["Histórico"]})
        if not self.hojas["Competencia"]:
            try:
                import crear_memoria
                import keyword_ml
                for c in crear_memoria.CAMPANAS:
                    for asin, motivo in c.get("asins", []):
                        self.hojas["Competencia"].append({
                            "Producto (ASIN)": keyword_ml.PERFILES[c["perfil"]]["asin"], "ASIN competidor": asin,
                            "Motivo": motivo, "Confirmado por Juan": "Pendiente", "Añadido": hoy.isoformat()})
            except Exception:  # sin crear_memoria la hoja queda vacía para que Juan la rellene
                pass

    # ------------------------------------------------------------ Diario y Seguimiento
    def guardar_diario(self, filas):
        """Upsert por (fecha, clave): los últimos días se vuelven a pedir porque Amazon
        sigue atribuyendo ventas hasta 7 días después del clic."""
        idx = {(str(f["Fecha"])[:10], str(f["Clave"])): i for i, f in enumerate(self.hojas["Diario"])}
        for f in filas:
            k = (str(f["Fecha"])[:10], str(f["Clave"]))
            if k in idx:
                self.hojas["Diario"][idx[k]] = f
            else:
                idx[k] = len(self.hojas["Diario"])
                self.hojas["Diario"].append(f)

    def guardar_foto(self, hoy, cuenta, acumulados, producto_de):
        """Una foto por elemento y día (si hay varias ejecuciones el mismo día, se queda la última)."""
        hoy_s = hoy.isoformat()
        self.hojas["Seguimiento"] = [f for f in self.hojas["Seguimiento"] if str(f["Fecha"])[:10] != hoy_s]
        for clave, el in cuenta.elementos.items():
            m = acumulados[clave]
            c = cuenta.campanas.get(el.id_campana)
            self.hojas["Seguimiento"].append({
                "Fecha": hoy_s, "Clave": clave, "ID campaña": el.id_campana, "Producto (ASIN)": producto_de(el),
                "Campaña": c.nombre if c else "", "Palabra clave / segmentación": el.texto,
                "Coincidencia": el.coincidencia, "Estado": el.estado, "Puja (€)": el.puja,
                "Impresiones": m.impresiones, "Clics": m.clics, "Coste (€)": round(m.coste, 2),
                "Compras": m.compras, "Ventas (€)": round(m.ventas, 2)})

    def series(self):
        return Series(self.hojas["Diario"], self.hojas["Seguimiento"])

    # ------------------------------------------------------------ Tickets
    def siguiente_ticket(self):
        n = max((int(str(t["Ticket"]).split("-")[-1]) for t in self.hojas["Tickets"] if t.get("Ticket")), default=0)
        return f"T-{n + 1:04d}"

    def registrar(self, cambio, hoy):
        tid = self.siguiente_ticket()
        self.hojas["Tickets"].append({
            "Ticket": tid, "Fecha": hoy.isoformat(), "Tipo": cambio.tipo, "Clave": cambio.clave,
            "Producto (ASIN)": cambio.producto, "ID campaña": cambio.id_campana, "Campaña": cambio.campana,
            "ID grupo": cambio.id_grupo, "Palabra clave / segmentación": cambio.texto,
            "Coincidencia": cambio.coincidencia, "Antes": cambio.antes, "Después": cambio.despues,
            "Motivo": cambio.motivo, "Estado": cambio.estado, "Detalle": cambio.detalle,
            "Requiere revisión": "Sí" if cambio.requiere_revision else "",
            "Base: clics": cambio.base.clics, "Base: coste": round(cambio.base.coste, 2),
            "Base: compras": cambio.base.compras, "Base: ventas": round(cambio.base.ventas, 2),
            "Datos extra": json.dumps(cambio.extra, ensure_ascii=False) if cambio.extra else None})
        return tid

    def tickets(self, clave=None, tipos=None, estados=("confirmado", "enviado_bulk")):
        out = [t for t in self.hojas["Tickets"]
               if (clave is None or str(t["Clave"]) == str(clave))
               and (tipos is None or t["Tipo"] in tipos)
               and (estados is None or t["Estado"] in estados)]
        return sorted(out, key=lambda t: (str(t["Fecha"]), str(t["Ticket"])))

    # ------------------------------------------------------------ Alertas
    def alerta_ya_enviada(self, tipo, clave, hoy):
        return any(str(a["Fecha"])[:10] == hoy.isoformat() and a["Tipo"] == tipo and str(a["Clave"]) == str(clave)
                   for a in self.hojas["Alertas"])

    def registrar_alerta(self, ahora, tipo, clave, mensaje, correo):
        self.hojas["Alertas"].append({"Fecha": ahora.date().isoformat(), "Hora": ahora.strftime("%H:%M"),
                                      "Tipo": tipo, "Clave": clave, "Mensaje": mensaje, "Correo": correo})

    # ------------------------------------------------------------ Competencia / Investigación
    def competidores(self, asin, solo_confirmados=True):
        return [c for c in self.hojas["Competencia"] if str(c["Producto (ASIN)"]) == asin and c.get("ASIN competidor")
                and (not solo_confirmados or str(c.get("Confirmado por Juan") or "").strip().lower() in ("sí", "si"))]

    def investigacion(self, asin):
        """Última investigación del producto: [(palabra, motivo)] en orden de rank."""
        filas = [f for f in self.hojas["Investigación"] if str(f["Producto (ASIN)"]) == asin]
        if not filas:
            return None, []
        ult = max(str(f["Fecha"])[:10] for f in filas)
        filas = sorted((f for f in filas if str(f["Fecha"])[:10] == ult), key=lambda f: num(f["Rank"], 999))
        return fecha(ult), filas


def _celda(v):
    if isinstance(v, float):
        return round(v, 4)
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def _m(f):
    return Metricas(num(f.get("Clics")), num(f.get("Coste (€)")), num(f.get("Compras")),
                    num(f.get("Ventas (€)")), num(f.get("Impresiones")))


class Series:
    """Acumulados por elemento y fecha, desde el Diario (API) o desde las fotos (Seguimiento).

    - Diario: filas por día -> el acumulado hasta una fecha es la suma de los días anteriores.
      Las ventas se atribuyen al día del clic, así que 'maduro' es exacto.
    - Fotos: acumulados tal como estaban el día de la foto. 'Maduro' = la foto de hace ≥ 7 días
      (sus clics tienen al menos 7 días). Sus ventas pueden estar incompletas, así que para
      "¿ha vendido algo?" se usan las ventas de la foto más reciente (una venta es una venta).
    """

    def __init__(self, diario, seguimiento):
        self.diario = defaultdict(list)
        self.fotos = defaultdict(list)
        self.campana = {}
        for f in diario:
            k = str(f["Clave"])
            self.diario[k].append((fecha(f["Fecha"]), _m(f)))
            if f.get("ID campaña"):
                self.campana[k] = str(f["ID campaña"])
        for f in seguimiento:
            k = str(f["Clave"])
            self.fotos[k].append((fecha(f["Fecha"]), _m(f)))
            if f.get("ID campaña"):
                self.campana.setdefault(k, str(f["ID campaña"]))
        self._pref = {}
        for k, filas in self.diario.items():
            filas.sort(key=lambda x: x[0])
            acum, pref = Metricas(), []
            for _, m in filas:
                acum = acum + m
                pref.append(acum)
            self._pref[k] = ([d for d, _ in filas], pref)
        for filas in self.fotos.values():
            filas.sort(key=lambda x: x[0])

    def tiene_diario(self, clave):
        return str(clave) in self._pref

    def acumulado(self, clave, hasta):
        """Acumulado del elemento con datos hasta 'hasta' (incluido)."""
        k = str(clave)
        if k in self._pref:
            fechas, pref = self._pref[k]
            i = bisect_right(fechas, hasta)
            return pref[i - 1] if i else Metricas()
        fotos = self.fotos.get(k, [])
        i = bisect_right([d for d, _ in fotos], hasta)
        return fotos[i - 1][1] if i else Metricas()

    def maduro(self, clave, hoy):
        """Clics y coste con más de 7 días. Compras/ventas: las de esos clics (Diario) o,
        con fotos, las conocidas hoy."""
        if self.tiene_diario(clave):
            return self.acumulado(clave, hoy - timedelta(days=config.DIAS_MADUREZ + 1))
        m = self.acumulado(clave, hoy - timedelta(days=config.DIAS_MADUREZ))
        hoy_m = self.acumulado(clave, hoy)
        return Metricas(m.clics, m.coste, hoy_m.compras, hoy_m.ventas, m.impresiones)

    def base_en(self, clave, fecha_cambio, base_guardada):
        """Acumulado en el momento de un cambio. Con Diario se recalcula (hasta el día anterior),
        así las ventas que llegaron tarde por clics de antes del cambio no se le atribuyen."""
        if self.tiene_diario(clave):
            return self.acumulado(clave, fecha_cambio - timedelta(days=1))
        return base_guardada

    def claves_de_campana(self, id_campana):
        return [k for k, c in self.campana.items() if c == str(id_campana)]

    def gasto_mes(self, hoy):
        """Gasto del mes en curso. None si no se puede saber (sin datos anteriores al mes)."""
        inicio = hoy.replace(day=1)
        if self._pref:
            return sum(m.coste for filas in self.diario.values() for d, m in filas if d >= inicio)
        claves = list(self.fotos)
        if not claves:
            return None
        antes = inicio - timedelta(days=1)
        if not any(f[0][0] <= antes for f in self.fotos.values() if f):
            return None
        return sum(self.acumulado(k, hoy).coste - self.acumulado(k, antes).coste for k in claves)
