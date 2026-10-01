"""
ficha.py — el agente de página de producto (la ficha de Amazon).

Python prepara, por producto, los datos con los que mejorar la ficha; Cowork (Claude, con navegador)
lee la ficha actual en amazon.es y escribe la propuesta (RUTINA_COWORK.md). La ficha en Seller
Central nunca se cambia sola: decide Juan.

salidas/<día>/ficha_datos_<día>.xlsx, una hoja por producto con cuatro bloques:
  1. Palabras sueltas: compras, clics y conversión de cada palabra en las búsquedas que ya se han
     pagado (histórico + campañas). Las que más venden van al título y a los bullets.
  2. Frases: las que venden (usarlas) y las que gastan sin vender (≥ 10 clics y 0 compras: no
     destacarlas, la ficha no convence a quien busca eso).
  3. Frases investigadas que pasan el filtro del agente (ACOS predicho ≤ 30 %): candidatas a
     términos de búsqueda ocultos (backend) y a bullets.
  4. Competencia: lo que dicen los compradores de los competidores (hoja Competencia y, si está,
     "Reseñas competencia"), para responder en la ficha a lo que critican.
Y una hoja "Propuesta" que rellena Cowork.

Uso suelto (lee la memoria más reciente):  python ficha.py
"""

from collections import defaultdict
from datetime import date
from pathlib import Path

import config
import keyword_ml as kml
from documento import num

MIN_CLICS_SIN_VENTA = 10


def _frases(doc, asin):
    """{firma: {frase, clics, compras, ventas, origen}} del histórico y de las campañas actuales."""
    out = {}

    def poner(frase, clics, compras, ventas, origen):
        fi = kml.firma(str(frase or ""))
        if not fi:
            return
        f = out.setdefault(fi, {"frase": str(frase).strip().lower(), "clics": 0.0, "compras": 0.0, "ventas": 0.0,
                                "origen": set()})
        f["clics"] += clics
        f["compras"] += compras
        f["ventas"] += ventas
        f["origen"].add(origen)

    for f in doc.hojas["Histórico"]:
        if str(f.get("Producto (ASIN)")) == asin:
            poner(f.get("Keyword / segmento"), num(f.get("Clics")), num(f.get("Compras")), num(f.get("Ventas (€)")),
                  "histórico")
    for f in doc.hojas["Segmentación"]:
        if str(f.get("Producto (ASIN)")) == asin and f.get("Tipo") == "keyword":
            poner(f.get("Palabra clave / segmentación"), num(f.get("Clics")), num(f.get("Compras")),
                  num(f.get("Ventas (€)")), "campañas")
    return out


def datos(doc, catalogo, asin):
    """Los cuatro bloques de un producto, como listas de filas."""
    frases = _frases(doc, asin)
    palabras = defaultdict(lambda: [0.0, 0.0])
    for f in frases.values():
        for w in set(kml.tokens_contenido(f["frase"])):
            palabras[w][0] += f["clics"]
            palabras[w][1] += f["compras"]
    sueltas = sorted(([w, c, k, round(k / c, 3) if c else None] for w, (c, k) in palabras.items() if c),
                     key=lambda r: (-r[2], -r[1]))

    venden = sorted((f for f in frases.values() if f["compras"] > 0), key=lambda f: -f["compras"])
    sin_venta = sorted((f for f in frases.values() if f["compras"] == 0 and f["clics"] >= MIN_CLICS_SIN_VENTA),
                       key=lambda f: -f["clics"])
    filas_frases = [[f["frase"], f["clics"], f["compras"], round(f["compras"] / f["clics"], 3) if f["clics"] else None,
                     ", ".join(sorted(f["origen"])), "usarla en título o bullets"] for f in venden]
    filas_frases += [[f["frase"], f["clics"], 0, 0.0, ", ".join(sorted(f["origen"])),
                      "no destacarla: la ficha no convence a quien busca esto"] for f in sin_venta]

    investigadas = []
    for f in doc.investigacion(asin)[1]:
        frase = str(f["Palabra clave"])
        if kml.firma(frase) in frases:
            continue
        ev = catalogo.evaluar(asin, frase, rec=num(f.get("Puja sugerida (€)"), None))
        if ev["pasa"]:
            investigadas.append([frase, round(ev["acos_pred"], 3), f.get("Volumen"), f.get("Fuente")])
    investigadas.sort(key=lambda r: r[1])

    corto = catalogo.producto(asin).corto.lower()
    competencia = [[c.get("ASIN competidor"), c.get("Motivo")] for c in doc.hojas["Competencia"]
                   if str(c.get("Producto (ASIN)")) == asin]
    resenas = doc.otras.get("Reseñas competencia") or []
    if resenas:
        cab = [str(x or "") for x in resenas[0]]
        for r in resenas[1:]:
            d = dict(zip(cab, r))
            if corto and corto in str(d.get("Producto nuestro") or "").lower():
                competencia.append([d.get("ASIN competidor"), " | ".join(
                    f"{k}: {d[k]}" for k in ("Competidor", "Lo que valoran", "Lo que critican", "Idea para keywords / ficha")
                    if d.get(k))])
    return sueltas, filas_frases, investigadas, competencia


def productos(doc, catalogo):
    """Los productos de FreshFinder: los que tienen ventas en el histórico o fila en la hoja Economía
    (no los ASIN ajenos que Amazon mete en sus campañas de recomendaciones)."""
    asins = {a for a, p in catalogo.productos.items() if p.compras > 0}
    asins |= {str(f["Producto (ASIN)"]) for f in doc.hojas["Economía"] if f.get("Producto (ASIN)")}
    return sorted(asins)


def escribir(doc, catalogo, ruta, hoy):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    negrita, azul = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    wb = Workbook()
    guia = wb.active
    guia.title = "Cómo usar"
    for linea in (
            f"Datos para mejorar las fichas de producto — {hoy:%d/%m/%Y}",
            "Una hoja por producto. Cowork lee la ficha actual en https://www.amazon.es/dp/<ASIN> y rellena 'Propuesta'.",
            "La ficha nunca se cambia sola: Juan decide qué aplicar en Seller Central.",
            "Palabras sueltas: las que más compras han traído van primero en el título y en los bullets.",
            "Frases 'no destacarla': han gastado ≥ 10 clics sin ninguna venta.",
            f"Investigadas: frases nuevas que pasan el filtro del agente (ACOS predicho ≤ {config.ACOS_MAX_KEYWORD_NUEVA:.0%}),"
            " buenas para los términos de búsqueda ocultos.",
            "Competencia: responde en la ficha a lo que critican de los competidores.",
            "Reglas de Amazon: bullets de 10 a 255 caracteres (mejor ≤ 200), sin exclamaciones ni 'el mejor'/'número 1'.",
            "Sin Registro de Marca: nada de contenido A+, Brand Store ni Sponsored Brands.",
            "Ofertas: un cupón o un descuento porcentual no afectan a la Oferta Destacada; el precio rebajado sí. Decide Juan."):
        guia.append([linea])
    guia["A1"].font = Font(bold=True, size=13)
    guia.column_dimensions["A"].width = 120

    def bloque(ws, titulo, cab, filas):
        ws.append([titulo])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        ws.append(cab)
        for c in ws[ws.max_row]:
            c.font, c.fill = negrita, azul
        for r in filas:
            ws.append(r)
        ws.append([])

    for asin in productos(doc, catalogo):
        sueltas, frases, investigadas, competencia = datos(doc, catalogo, asin)
        p = catalogo.producto(asin)
        ws = wb.create_sheet(f"{p.corto} ({asin})"[:31])
        ws.append([f"{p.corto} — {asin} — https://www.amazon.es/dp/{asin}"])
        ws.append([p.nombre])
        ws.append([])
        bloque(ws, "1. Palabras sueltas", ["Palabra", "Clics", "Compras", "Conversión"], sueltas[:40])
        bloque(ws, "2. Frases que venden y frases que gastan sin vender",
               ["Frase", "Clics", "Compras", "Conversión", "De dónde", "En la ficha"], frases)
        bloque(ws, "3. Frases investigadas que pasan el filtro", ["Frase", "ACOS predicho", "Volumen", "Fuente"],
               investigadas[:40])
        bloque(ws, "4. Competencia", ["ASIN competidor", "Qué dicen los compradores"], competencia)
        ws.column_dimensions["A"].width = 50
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["F"].width = 45

    prop = wb.create_sheet("Propuesta")
    prop.append(["Producto (ASIN)", "Elemento", "Texto actual", "Texto propuesto", "Por qué (datos que lo apoyan)"])
    for c in prop[1]:
        c.font, c.fill = negrita, azul
    for asin in productos(doc, catalogo):
        for elemento in ("Título", "Bullet 1", "Bullet 2", "Bullet 3", "Bullet 4", "Bullet 5",
                         "Términos de búsqueda ocultos", "Imágenes", "Oferta / cupón"):
            prop.append([asin, elemento])
    for col, ancho in zip("ABCDE", (16, 28, 60, 60, 60)):
        prop.column_dimensions[col].width = ancho
    Path(ruta).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


def main():
    import carpetas
    from documento import Documento
    from prediccion import Catalogo
    hoy = date.today()
    memoria = carpetas.memoria_anterior(hoy)
    if not memoria:
        raise SystemExit(f"No hay ninguna memoria en {config.SALIDAS}/: ejecuta antes el agente.")
    ruta = escribir(Documento(memoria), Catalogo(), carpetas.salida(hoy) / f"ficha_datos_{hoy.isoformat()}.xlsx", hoy)
    print(f"Datos de las fichas (de {memoria}) -> {ruta}")


if __name__ == "__main__":
    main()
