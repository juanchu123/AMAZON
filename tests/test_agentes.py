"""Agente de finanzas (finanzas.py) y agente de página de producto (ficha.py)."""

import sys
from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
import ficha
import finanzas
import pujas
from documento import Documento
from modelo import ACTIVO, SOLO_BAJA, Campana
from prediccion import Catalogo

RAIZ = Path(__file__).resolve().parent.parent
PINZA = "B0DCZS1NR6"
HOY = date(2026, 10, 20)


@pytest.fixture(scope="module")
def catalogo():
    return Catalogo()


@pytest.fixture(autouse=True)
def equilibrios_limpios(monkeypatch):
    monkeypatch.setattr(config, "ACOS_EQUILIBRIO_POR_ASIN", {})


def test_margen_y_acos_de_equilibrio():
    # pinza: 11,24 € con IVA, 21 % IVA, 15 % comisión, 3 € FBA, 0,80 € de coste
    m = finanzas.margen_unitario(11.24, 0.21, 0.15, 3.0, 0.80)
    assert m == pytest.approx(11.24 / 1.21 - 11.24 * 0.15 - 3.8)
    assert m / 11.24 == pytest.approx(0.338, abs=0.001)


def test_economia_se_siembra_sin_pisar_a_juan_y_llega_a_las_pujas(tmp_path, catalogo):
    doc = Documento(tmp_path / "m.xlsx")
    finanzas.sembrar(doc, catalogo)
    fila = next(f for f in doc.hojas["Economía"] if f["Producto (ASIN)"] == PINZA)
    assert fila["Ticket medio con IVA (€)"] == pytest.approx(11.24, abs=0.01) and fila["Comisión Amazon"] is None
    assert finanzas.aplicar(doc) == {} and config.acos_equilibrio(PINZA) == config.ACOS_EQUILIBRIO_DEFECTO

    fila.update({"Comisión Amazon": "15 %", "Tarifa FBA (€/ud)": 3, "Coste del producto (€/ud)": "0,80"})
    finanzas.sembrar(doc, catalogo)                                        # no pisa lo escrito
    assert fila["Comisión Amazon"] == "15 %"
    eq = finanzas.aplicar(doc)[PINZA]
    assert eq == pytest.approx(0.338, abs=0.001) and config.acos_equilibrio(PINZA) == eq
    assert finanzas.avisos(doc) and "34%" in finanzas.avisos(doc)[0][2]    # por debajo del objetivo 30-35 %

    # el tope de las pujas usa ese equilibrio: el clic más caro nunca pasa de p × ticket × 33,8 %
    c = Campana("C", "c", ACTIVO, 7, estrategia_pujas=SOLO_BAJA)
    calc = pujas.calcular(c, 0.10, 11.24, 1.0, PINZA)                       # agresividad 1 -> ACOS objetivo 35 %
    assert calc.tope == pytest.approx(0.10 * 11.24 * eq)
    assert calc.puja <= calc.tope + 1e-9 < calc.objetivo


def test_informe_de_beneficio_por_producto(tmp_path, catalogo):
    doc = Documento(tmp_path / "m.xlsx")
    doc.hojas["Economía"].append({"Producto (ASIN)": PINZA, "Producto": "Pinza", "Ticket medio con IVA (€)": 11.24,
                                  "IVA": 0.21, "Comisión Amazon": 0.15, "Tarifa FBA (€/ud)": 3,
                                  "Coste del producto (€/ud)": 0.8})
    finanzas.aplicar(doc)
    for d, clics, coste, compras in ((HOY - timedelta(days=40), 100, 50.0, 5), (HOY, 160, 80.0, 9)):
        doc.hojas["Seguimiento"].append({"Fecha": d.isoformat(), "Clave": "K1", "ID campaña": "C", "Producto (ASIN)": PINZA,
                                         "Clics": clics, "Coste (€)": coste, "Compras": compras,
                                         "Ventas (€)": compras * 11.24})
    filas = finanzas.informe(doc, doc.series(), catalogo, HOY)
    f30 = next(f for f in filas if f["Producto (ASIN)"] == PINZA and f["Periodo"] == "Últimos 30 días")
    margen = finanzas.margen_unitario(11.24, 0.21, 0.15, 3, 0.8)
    assert f30["Gasto en anuncios (€)"] == 30.0 and f30["Compras por anuncios"] == 4
    assert f30["Beneficio después de publicidad (€)"] == pytest.approx(4 * round(margen, 2) - 30)   # margen al céntimo
    assert f30["Situación"] == "pierde dinero"
    assert filas[-1]["Producto (ASIN)"] == "TOTAL"


def test_datos_de_ficha_con_la_memoria_real(tmp_path, catalogo):
    doc = Documento(RAIZ / "salidas" / "2026-09-28" / config.NOMBRE_MEMORIA)
    sueltas, frases, investigadas, competencia = ficha.datos(doc, catalogo, PINZA)
    assert sueltas and sueltas[0][2] >= sueltas[-1][2]                     # ordenadas por compras
    assert any(f[5].startswith("usarla") for f in frases)
    assert all(f[1] <= config.ACOS_MAX_KEYWORD_NUEVA for f in investigadas)
    assert competencia                                                      # hoja Competencia de la memoria
    ruta = ficha.escribir(doc, catalogo, tmp_path / "ficha.xlsx", HOY)
    wb = load_workbook(ruta, read_only=True)
    assert "Propuesta" in wb.sheetnames and any(PINZA in h for h in wb.sheetnames)
