"""Sistema de pujas (pujas.py) con la hoja masiva REAL recortada (tests/fixtures/hoja_masiva_recortada.xlsx).

La hoja recortada conserva las cabeceras reales de Amazon y las campañas "Soporte móvil AC" (al alza y a
la baja) y "Todo" (terminada el 21/01/2025). A "Soporte móvil AC" se le ha quitado la fecha de
finalización para poder probarla activa, y sus keywords llevan métricas inventadas (la descarga real
del 28/09/2026 trae todo a 0).
"""

import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import agente
import config
import fuente_bulk
import pujas
from documento import Documento
from modelo import (ACTIVO, ALZA_BAJA, ESTRATEGIA, FINALIZADA, NUEVA_KEYWORD, PUJA, PUJA_FIJA, SOLO_BAJA, Campana,
                    Metricas)
from prediccion import Catalogo

FIXTURE = Path(__file__).parent / "fixtures" / "hoja_masiva_recortada.xlsx"
AC, TODO = "562775822100105", "319935329473301"
HOY = date(2026, 10, 20)


@pytest.fixture(scope="module")
def catalogo():
    return Catalogo()


@pytest.fixture(autouse=True)
def sin_correo_real(tmp_path, monkeypatch):
    for v in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "ANTHROPIC_API_KEY"):
        monkeypatch.delenv(v, raising=False)
    monkeypatch.setattr(config, "CORREOS_PENDIENTES", tmp_path / "correos")


def leer(ruta=FIXTURE, hoy=HOY):
    return fuente_bulk.FuenteBulk(ruta).leer_cuenta(hoy)


def campana(estrategia, **ajustes):
    return Campana("X", "x", ACTIVO, 10, estrategia_pujas=estrategia, ajustes_emplazamiento=ajustes)


# ---------------------------------------------------------------- lectura de la hoja real
def test_lee_estrategia_emplazamientos_fin_y_coste():
    cuenta = leer()
    ac, todo = cuenta.campanas[AC], cuenta.campanas[TODO]
    assert ac.estado == ACTIVO and ac.estrategia_pujas == ALZA_BAJA
    assert set(ac.ajustes_emplazamiento) == {"superior", "resto de la búsqueda", "página del producto", "Amazon Business"}
    assert todo.estado == FINALIZADA and todo.fecha_fin == date(2025, 1, 21)
    # el coste sale de la columna "Inversión"
    assert sum(e.metricas.coste for e in cuenta.elementos.values()) == pytest.approx(30 + 27 + 12 + 6)


# ---------------------------------------------------------------- multiplicador (ayuda de Amazon Ads)
def test_multiplicador_segun_estrategia_y_emplazamiento():
    assert pujas.multiplicador(campana(SOLO_BAJA), SOLO_BAJA) == 1
    assert pujas.multiplicador(campana(PUJA_FIJA), PUJA_FIJA) == 1
    assert pujas.multiplicador(campana(ALZA_BAJA), ALZA_BAJA) == 2              # +100 % arriba de la búsqueda
    # ejemplo de Amazon: +50 % arriba y al alza y a la baja -> hasta ×3 arriba
    assert pujas.multiplicador(campana(ALZA_BAJA, superior=50), ALZA_BAJA) == pytest.approx(3.0)
    assert pujas.multiplicador(campana(SOLO_BAJA, superior=20), SOLO_BAJA) == pytest.approx(1.2)
    assert pujas.multiplicador(campana(""), "") == 2                            # desconocida: el peor caso


def test_puja_efectiva_max_nunca_pasa_del_equilibrio_en_las_campanas_reales():
    cuenta = leer()
    for id_c in (AC, TODO):
        c = cuenta.campanas[id_c]                    # sin elegir estrategia: la que tienen, al alza y a la baja
        for p in (0.01, 0.047, 0.11, 0.3):
            for aggr in (0.0, 0.5, 1.0):
                calc = pujas.calcular(c, p, 14.27, aggr)
                assert pujas.puja_efectiva_max(c, calc.puja) <= calc.tope + 0.005 or calc.puja == config.PUJA_MINIMA_AMAZON


def test_al_alza_y_a_la_baja_divide_la_puja_entre_2_y_lo_explica():
    c = campana(ALZA_BAJA)
    calc = pujas.calcular(c, 0.10, 14.0, 0.5)
    assert calc.multiplicador == 2
    assert calc.puja == round(calc.tope / 2, 2) < round(calc.objetivo, 2)
    assert "/ 2.00" in calc.explicacion(0.10, 14.0)


# ---------------------------------------------------------------- el agente elige la estrategia
def test_sin_datos_elige_solo_a_la_baja():
    est, motivo = pujas.elegir_estrategia(campana(ALZA_BAJA), Metricas())
    assert est == SOLO_BAJA and "compras" in motivo


def test_probada_pero_sin_margen_elige_solo_a_la_baja():
    m = Metricas(clics=200, coste=30, compras=12, ventas=171)              # ACOS 17,5 %
    est, motivo = pujas.elegir_estrategia(campana(SOLO_BAJA), m)
    assert est == SOLO_BAJA and "no compensa" in motivo                    # 35 % × 2 > 35 %


def test_probada_y_con_margen_elige_al_alza_y_no_recorta_la_puja(monkeypatch):
    monkeypatch.setattr(config, "ACOS_EQUILIBRIO_DEFECTO", 0.80)          # p. ej. con la hoja Economía
    m = Metricas(clics=200, coste=30, compras=12, ventas=171)
    c = campana(SOLO_BAJA)
    est, _ = pujas.elegir_estrategia(c, m)
    assert est == ALZA_BAJA
    c.estrategia_objetivo = est
    calc = pujas.calcular(c, 0.10, 14.0, 0.5)
    assert calc.puja == round(calc.objetivo, 2)                            # 32,5 % × 2 ≤ 80 %: no hace falta recortar
    assert pujas.puja_efectiva_max(c, calc.puja) <= calc.tope + 0.005


def test_probada_con_acos_malo_no_sube():
    m = Metricas(clics=400, coste=150, compras=12, ventas=171)             # ACOS 88 %
    assert pujas.elegir_estrategia(campana(ALZA_BAJA), m)[0] == SOLO_BAJA


# ---------------------------------------------------------------- una ronda con la hoja real
def _doc_con_foto(tmp_path, cuenta):
    """Una foto de hace 10 días con los mismos acumulados: los clics ya son maduros."""
    doc = Documento(tmp_path / "doc.xlsx")
    for e in cuenta.elementos.values():
        m = e.metricas
        doc.hojas["Seguimiento"].append({"Fecha": (HOY - timedelta(days=10)).isoformat(), "Clave": e.clave,
                                         "ID campaña": e.id_campana, "Clics": m.clics, "Coste (€)": m.coste,
                                         "Compras": m.compras, "Ventas (€)": m.ventas})
    doc.guardar()
    return Documento(tmp_path / "doc.xlsx")


def _ronda(tmp_path, ruta, doc, catalogo, hoy):
    fuente = fuente_bulk.FuenteBulk(ruta, salida_dir=tmp_path)
    return agente.ejecutar(fuente, doc, catalogo, hoy, datetime.combine(hoy, datetime.min.time()), investigar="no",
                           log=lambda *_: None)


def test_ronda_real_cambia_estrategia_y_pujas_dentro_del_tope(tmp_path, catalogo):
    doc = _doc_con_foto(tmp_path, leer())
    res = _ronda(tmp_path, FIXTURE, doc, catalogo, HOY)
    cambios = res["cambios"]

    est = [c for c in cambios if c.tipo == ESTRATEGIA]
    assert [(c.id_campana, c.antes, c.despues, c.estado) for c in est] == [(AC, ALZA_BAJA, SOLO_BAJA, "enviado_bulk")]
    assert not [c for c in cambios if c.id_campana == TODO]                # terminada: no se toca

    cuenta = leer()
    cuenta.campanas[AC].estrategia_objetivo = SOLO_BAJA
    pujadas = [c for c in cambios if c.tipo in (PUJA, NUEVA_KEYWORD)]
    assert pujadas, "con 60 clics y 6 compras tiene que mover alguna puja"
    for c in pujadas:
        assert c.despues * pujas.multiplicador(cuenta.campanas[AC]) <= c.extra["tope_rentable"] + 0.005

    # la hoja de cambios lleva la estrategia en la fila de la campaña
    ws = load_workbook(res["bulk"])[fuente_bulk.HOJA]
    cab = [x.value for x in ws[1]]
    filas = [dict(zip(cab, [x.value for x in r])) for r in ws.iter_rows(min_row=2)]
    assert any(f["Entidad"] == "Campaña" and str(f["ID de la campaña"]) == AC and f["Estrategia de pujas"] == SOLO_BAJA
               for f in filas)

    # siguiente descarga: Juan subió la hoja -> estrategia confirmada y no se vuelve a pedir
    wb = load_workbook(FIXTURE)
    ws = wb[fuente_bulk.HOJA]
    cab = [x.value for x in ws[1]]
    i_est, i_c = cab.index("Estrategia de pujas"), cab.index("ID de la campaña")
    for r in ws.iter_rows(min_row=2):
        if str(r[i_c].value) == AC and r[i_est].value:
            r[i_est].value = SOLO_BAJA
    wb.save(tmp_path / "descarga2.xlsx")
    res2 = _ronda(tmp_path, tmp_path / "descarga2.xlsx", Documento(tmp_path / "doc.xlsx"), catalogo, HOY + timedelta(days=1))
    assert not [c for c in res2["cambios"] if c.tipo == ESTRATEGIA]
    t = next(t for t in Documento(tmp_path / "doc.xlsx").hojas["Tickets"] if t["Tipo"] == ESTRATEGIA)
    assert t["Estado"] == "confirmado"


def test_todo_terminado_para_el_agente(tmp_path, catalogo):
    """La descarga real del 28/09: las 2 campañas 'activadas' ya terminaron. El agente para y avisa."""
    doc = Documento(tmp_path / "doc.xlsx")
    res = _ronda(tmp_path, Path(__file__).parent.parent / "datos" / "hoja_masiva_2026-09-28.xlsx", doc, catalogo,
                 date(2026, 9, 28))
    assert res["parado"] and "terminadas" in res["parado"]
