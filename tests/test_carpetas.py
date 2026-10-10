"""entradas/ y salidas/ con una carpeta por día; la memoria parte de la del día anterior."""

import shutil
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import agente
import carpetas
import config
import investigacion
from documento import Documento

RAIZ = Path(__file__).resolve().parent.parent
DIA = RAIZ / "entradas" / "2026-09-28"


@pytest.fixture
def carpetas_tmp(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "ENTRADAS", tmp_path / "entradas")
    monkeypatch.setattr(config, "SALIDAS", tmp_path / "salidas")
    for v in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "ANTHROPIC_API_KEY", "AMAZON_ADS_CLIENT_ID"):
        monkeypatch.delenv(v, raising=False)
    return tmp_path


def test_entrada_mas_reciente_con_hoja_masiva(carpetas_tmp):
    for d in ("2026-09-20", "2026-09-28", "2026-10-02"):
        (config.ENTRADAS / d).mkdir(parents=True)
    shutil.copy(DIA / "hoja_masiva_2026-09-28.xlsx", config.ENTRADAS / "2026-09-20")
    shutil.copy(DIA / "hoja_masiva_2026-09-28.xlsx", config.ENTRADAS / "2026-09-28")
    (config.ENTRADAS / "no-es-un-dia").mkdir()
    assert carpetas.entrada(date(2026, 10, 5)).name == "2026-09-28"     # la del 02/10 no trae hoja masiva
    assert carpetas.entrada(date(2026, 9, 25)).name == "2026-09-20"     # nunca una carpeta del futuro


def test_memoria_parte_de_la_anterior_y_conserva_hojas_ajenas(carpetas_tmp):
    ayer = config.SALIDAS / "2026-09-28"
    ayer.mkdir(parents=True)
    shutil.copy(RAIZ / "salidas" / "2026-09-28" / config.NOMBRE_MEMORIA, ayer)
    origen = carpetas.memoria_anterior(date(2026, 9, 30))
    assert origen == ayer / config.NOMBRE_MEMORIA
    doc = Documento(carpetas.salida(date(2026, 9, 30)) / config.NOMBRE_MEMORIA, origen=origen)
    tickets = len(doc.hojas["Tickets"])
    assert tickets >= 31 and "Productos" in doc.otras
    doc.guardar()
    wb = load_workbook(config.SALIDAS / "2026-09-30" / config.NOMBRE_MEMORIA, read_only=True)
    assert {"Productos", "Anuncios (histórico)", "Tickets"} <= set(wb.sheetnames)
    assert len(Documento(config.SALIDAS / "2026-09-30" / config.NOMBRE_MEMORIA).hojas["Tickets"]) == tickets
    assert (ayer / config.NOMBRE_MEMORIA).exists()                        # la del día anterior no se toca


CSV = ("ASIN;Palabra clave;Motivo;Fuente;Volumen;Puja sugerida (€)\n"
       "B0DCZS1NR6;soporte movil coche pinza fuerte;autocompletado;autocompletado;alto;0,31\n"
       "B0DCZS1NR6;Soporte móvil PINZA coche fuerte;otra vez en otro orden;autocompletado;alto;\n"
       "B0DHYBY6MS;soporte telefono rejilla ventilacion coche;autocompletado;autocompletado;alto;\n")


def test_csv_de_cowork_sin_repetir(tmp_path):
    (tmp_path / "investigacion_2026-09-29.csv").write_text(CSV, encoding="utf-8")
    doc = Documento(tmp_path / "m.xlsx")
    res = investigacion.importar_entrada(doc, tmp_path, date(2026, 9, 29))
    assert {(a, n) for _, a, n in res} == {("B0DCZS1NR6", 1), ("B0DHYBY6MS", 1)}   # la 2ª frase es la misma
    fila = doc.hojas["Investigación"][0]
    assert fila["Puja sugerida (€)"] == 0.31 and fila["Fecha"] == "2026-09-29"
    # al día siguiente, la misma frase no vuelve a entrar; una nueva sí
    (tmp_path / "investigacion_2026-09-29.csv").write_text(CSV + "B0DCZS1NR6;pinza movil salpicadero curvo;x;x;;\n",
                                                          encoding="utf-8")
    res = investigacion.importar_entrada(doc, tmp_path, date(2026, 9, 30))
    assert ("B0DCZS1NR6", 1) in {(a, n) for _, a, n in res} and len(doc.hojas["Investigación"]) == 3
    # el agente ve todas las frases de todos los días, una sola vez
    assert len(doc.investigacion("B0DCZS1NR6")[1]) == 2


def test_csv_con_cabecera_mala_avisa(tmp_path):
    (tmp_path / "investigacion_x.csv").write_text("frase;asin\nhola;B0DCZS1NR6\n", encoding="utf-8")
    with pytest.raises(SystemExit, match="faltan las columnas"):
        investigacion.importar_entrada(Documento(tmp_path / "m.xlsx"), tmp_path, date(2026, 9, 29))


def test_ya_vistas_para_cowork(carpetas_tmp):
    (config.SALIDAS / "2026-09-28").mkdir(parents=True)
    shutil.copy(RAIZ / "salidas" / "2026-09-28" / config.NOMBRE_MEMORIA, config.SALIDAS / "2026-09-28")
    investigacion.main(["--ya-vistas", "--hoy", "2026-09-29"])
    lineas = (config.SALIDAS / "2026-09-29" / "frases_ya_vistas.csv").read_text(encoding="utf-8-sig").splitlines()
    assert lineas[0] == "ASIN;Palabra clave;Firma;Origen" and len(lineas) > 100
    firmas = [tuple(l.split(";")[::2]) for l in lineas[1:]]
    assert len(firmas) == len(set(firmas))                                  # cada frase una vez por producto


def test_importa_el_excel_de_investigacion_sin_ml(tmp_path):
    doc = Documento(tmp_path / "m.xlsx")
    n = investigacion.importar_documento(doc, RAIZ / "salidas" / "2026-09-28" / "Documento_investigacion_keywords.xlsx",
                                         date(2026, 9, 28))
    assert n["B0DCZS1NR6"] > 10 and n["B0DHYBY6MS"] > 10
    assert not [f for f in doc.hojas["Investigación"] if f["Fuente"].startswith("Recomendada ML")]


def test_ejecucion_completa_escribe_en_salidas_del_dia(carpetas_tmp):
    (config.ENTRADAS / "2026-09-28").mkdir(parents=True)
    shutil.copy(DIA / "hoja_masiva_2026-09-28.xlsx", config.ENTRADAS / "2026-09-28")
    (config.ENTRADAS / "2026-09-28" / "investigacion_2026-09-28.csv").write_text(CSV, encoding="utf-8")
    (config.ENTRADAS / "2026-09-29").mkdir()                              # día sin hoja masiva, solo investigación
    (config.ENTRADAS / "2026-09-29" / "investigacion_2026-09-29.csv").write_text(
        "ASIN;Palabra clave;Motivo;Fuente;Volumen;Puja sugerida (€)\nB0DCZS1NR6;pinza movil salpicadero curvo;x;x;;\n",
        encoding="utf-8")
    (config.SALIDAS / "2026-09-28").mkdir(parents=True)
    shutil.copy(RAIZ / "salidas" / "2026-09-28" / config.NOMBRE_MEMORIA, config.SALIDAS / "2026-09-28")
    with pytest.raises(SystemExit) as fin:                                # descarga real: todo terminado -> para
        agente.main(["--simular", "--sin-investigacion", "--hoy", "2026-09-29"])
    assert fin.value.code == 2
    hoy = config.SALIDAS / "2026-09-29"
    doc = Documento(hoy / config.NOMBRE_MEMORIA)
    resumen = {f["Concepto"]: f["Valor"] for f in doc.hojas["Resumen"]}
    assert resumen["Memoria de partida"].endswith("2026-09-28/memoria_agente.xlsx")
    # las frases del CSV del 28 ya estaban en la memoria real (Cowork las importó ese día): no se repiten;
    # la del 29 (carpeta sin hoja masiva) sí entra
    assert not any("investigacion_2026-09-28.csv" in k for k in resumen)
    assert resumen["Investigación B0DCZS1NR6 (investigacion_2026-09-29.csv): frases nuevas"] == 1
    excel = load_workbook(hoy / "Documento_investigacion_keywords.xlsx", read_only=True)
    assert "Resumen" in excel.sheetnames and any("B0DCZS1NR6" in h for h in excel.sheetnames)
    assert len(doc.hojas["Tickets"]) >= 31 and "Productos" in doc.otras
    assert list((hoy / "correos_pendientes").glob("*.eml"))
