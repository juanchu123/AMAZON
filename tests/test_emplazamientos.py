"""Ajustes de emplazamiento gestionados por el agente (Juan, 30/09/2026)."""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
import emplazamientos
import pujas
from modelo import (ACTIVO, EMPLAZAMIENTO, PAGINA_PRODUCTO, RESTO_BUSQUEDA, SOLO_BAJA, SUPERIOR, Campana, Cuenta,
                    Metricas)
from test_agente import HOY


def _campana(top=(100, 10), resto=(200, 6), pagina=(100, 2), ajustes=None):
    c = Campana("C1", "Pinza", ACTIVO, 10.0, estrategia_pujas=SOLO_BAJA, ajustes_emplazamiento=ajustes or {})
    for lugar, (clics, compras) in ((SUPERIOR, top), (RESTO_BUSQUEDA, resto), (PAGINA_PRODUCTO, pagina)):
        c.metricas_emplazamiento[lugar] = Metricas(clics, clics * 0.4, compras, compras * 11.24)
    return c


def test_sube_donde_convierte_mejor_y_el_peor_queda_a_cero():
    cuenta = Cuenta(fecha=HOY)
    cuenta.campanas["C1"] = _campana()
    cambios = emplazamientos.decidir(cuenta)
    assert len(cambios) == 1 and cambios[0].tipo == EMPLAZAMIENTO
    nuevos = cambios[0].despues
    assert nuevos[PAGINA_PRODUCTO] == 0                       # convierte peor: base
    assert nuevos[SUPERIOR] > nuevos[RESTO_BUSQUEDA] > 0       # 10 % > 3 % > 2 %


def test_sin_datos_no_se_toca():
    cuenta = Cuenta(fecha=HOY)
    cuenta.campanas["C1"] = _campana(top=(10, 1), resto=(10, 1), pagina=(10, 0))
    assert emplazamientos.decidir(cuenta) == []


def test_ningun_clic_pasa_del_equilibrio_en_ningun_emplazamiento():
    cuenta = Cuenta(fecha=HOY)
    c = _campana()
    cuenta.campanas["C1"] = c
    emplazamientos.decidir(cuenta)
    p, ticket = 0.06, 11.24
    calc = pujas.calcular(c, p, ticket, 1.0)
    ajustes = pujas.ajustes_gestionados(c)
    for lugar in (SUPERIOR, RESTO_BUSQUEDA, PAGINA_PRODUCTO):
        coste_clic = calc.puja * (1 + ajustes[lugar] / 100)
        vale = p * c.ratios_emplazamiento[lugar] * ticket * config.acos_equilibrio()
        assert coste_clic <= vale + 0.005, lugar
    # y la puja base es para el peor emplazamiento (página de producto)
    assert calc.objetivo == pytest.approx(p * ticket * config.ACOS_OBJETIVO_MAX * c.ratios_emplazamiento[PAGINA_PRODUCTO], rel=0.01)


def test_hoja_masiva_escribe_ajuste_de_puja(tmp_path):
    from openpyxl import load_workbook
    import fuente_bulk
    from test_agente import cuenta_pinza
    cuenta = cuenta_pinza()
    cuenta.campanas["C1"] = _campana()
    cuenta.campanas["C1"].nombre = "Pinza - Principal V2"
    cambios = emplazamientos.decidir(cuenta)
    f = fuente_bulk.FuenteBulk(tmp_path / "x.xlsx", tmp_path)
    f.aplicar(cambios[0], cuenta)
    ws = load_workbook(f.cerrar(HOY))[fuente_bulk.HOJA]
    cab = [c.value for c in ws[1]]
    filas = [dict(zip(cab, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    assert {r["Emplazamiento"] for r in filas} == {"Emplazamiento superior", "Emplazamiento del resto de la búsqueda"}
    assert all(r["Entidad"] == "Ajuste de puja" and r["Operación"] == "Actualizar" for r in filas)
